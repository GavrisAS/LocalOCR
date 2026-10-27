"""Hard parser/writer deadlines and useful, literal linear salvage contracts."""

import json
import os
import sys
import time
from pathlib import Path
from threading import Event, Timer

import pytest
from openpyxl import load_workbook
from PIL import Image

from localocr import failsoft, postprocess, service
from localocr.failsoft import Extraction, excel_text, extract, write_extraction
from localocr.ollama import Cancelled, OCRServerError
from localocr.providers import OpenWebUIChatProvider, OpenWebUISettings
from localocr.supervisor import launch_owned, terminate_owned


class Provider(OpenWebUIChatProvider):
    def doctor(self):
        self.model_identity = {"digest": "synthetic-fixed"}
        return "ready"

    def request(self, method, endpoint, payload=None):
        self.http_ocr_calls = getattr(self, "http_ocr_calls", 0) + 1
        return {
            "choices": [{"message": {"content": "наблюдение 001\n=1+1"}, "finish_reason": "stop"}]
        }


def test_ten_thousand_rows_are_sequential_and_exact(tmp_path, monkeypatch):
    # Any hidden max_row scan of the growing salvage sheet fails immediately.
    from openpyxl.worksheet.worksheet import Worksheet

    original = Worksheet.max_row.fget

    def checked(sheet):
        if sheet.title == "Распознанный текст" and len(sheet._cells) > 1:
            raise AssertionError("writer scanned max_row")
        return original(sheet)

    # Limit the instrumentation to writing; readback needs dimensions.
    original_save = failsoft.Workbook.save

    def saved(book, path):
        monkeypatch.setattr(Worksheet, "max_row", property(original))
        return original_save(book, path)

    monkeypatch.setattr(Worksheet, "max_row", property(checked))
    monkeypatch.setattr(failsoft.Workbook, "save", saved)
    lines = [f"  Русское наблюдение {i:05} \\ =1+1  " for i in range(10000)]
    state = write_extraction(Extraction(lines=lines), tmp_path / "out.xlsx")
    book = load_workbook(tmp_path / "out.xlsx", read_only=True)
    assert [r[0] for r in book["Распознанный текст"].values] == [excel_text(v) for v in lines]
    assert state["nonempty_data_cells"] == 10000
    book.close()


@pytest.mark.parametrize(
    "text",
    [
        "=1+1",
        "+1",
        "-0",
        "@A",
        "a\x00b",
        "\\u0000",
        "x\ud800y",
        "a\ufffeb\uffff",
        "😀" * 20000,
        "  Пробелы  ",
    ],
)
def test_xml_literal_roundtrip(tmp_path, text):
    ex = Extraction(lines=[text])
    write_extraction(ex, tmp_path / "out.xlsx")
    book = load_workbook(tmp_path / "out.xlsx")
    assert "".join(r[0].value for r in book["Распознанный текст"]) == excel_text(text)
    assert all(r[0].data_type == "s" for r in book["Распознанный текст"])
    assert all(
        len(r[0].value.encode("utf-16-le")) // 2 <= 32767 for r in book["Распознанный текст"]
    )
    book.close()


@pytest.mark.parametrize("limit", ["MAX_XLSX_ROWS", "MAX_SALVAGE_ROWS"])
def test_row_limit_before_any_workbook_allocation(tmp_path, monkeypatch, limit):
    monkeypatch.setattr(failsoft, limit, 2)
    monkeypatch.setattr(failsoft, "Workbook", lambda: pytest.fail("allocated before cap"))
    with pytest.raises(ValueError, match="EXCEL_ROW_LIMIT"):
        write_extraction(Extraction(lines=["a", "b", "c"]), tmp_path / "out.xlsx")
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize(
    "mode,code",
    [
        ("parser_timeout", "PARSER_TIMEOUT"),
        ("writer_timeout", "WRITER_TIMEOUT"),
        ("crash", "RESOURCE_FAILURE"),
        ("resource", "RESOURCE_FAILURE"),
        ("io", "IO_FAILURE"),
        ("corrupt", "IO_FAILURE"),
    ],
)
def test_hard_fault_then_next_image(tmp_path, monkeypatch, mode, code):
    provider = Provider(OpenWebUISettings(api="chat", retries=0))
    monkeypatch.setattr(service, "create_provider", lambda *a, **kw: provider)
    folder = tmp_path / "вход"
    folder.mkdir()
    for name in ["а.png", "б.png"]:
        Image.new("RGB", (8, 8)).save(folder / name)
    original = postprocess.worker_argv
    count = 0

    def argv(spool, staging):
        nonlocal count
        count += 1
        if count == 1:
            return [
                sys.executable,
                str(Path("tests/helpers/stage27_fault_worker.py").resolve()),
                mode,
                str(spool),
                str(staging),
            ]
        return original(spool, staging)

    monkeypatch.setattr(postprocess, "worker_argv", argv)
    monkeypatch.setattr(postprocess, "PARSER_SECONDS", 1.0)
    monkeypatch.setattr(postprocess, "WRITER_SECONDS", 0.5)
    results = service.convert(folder, provider.settings, tmp_path / "выход", fail_soft=True)
    assert results[0].error_codes == [code]
    assert results[0].status == "FAILED" and results[0].output is None
    assert results[1].xlsx_with_data and results[1].status == "PARTIAL"
    assert not list((tmp_path / "выход").glob(".localocr-*"))
    assert not list((tmp_path / "runtime/work/postprocess").iterdir())


def test_owned_writer_cancel_never_publishes(tmp_path, monkeypatch):
    monkeypatch.setattr(
        postprocess,
        "worker_argv",
        lambda s, t: [
            sys.executable,
            str(Path("tests/helpers/stage27_fault_worker.py").resolve()),
            "writer_timeout",
            str(s),
            str(t),
        ],
    )
    cancel = Event()
    timer = Timer(0.5, cancel.set)
    timer.start()
    try:
        with pytest.raises(Cancelled):
            postprocess.convert_raw("наблюдение", tmp_path / "out.xlsx", cancel=cancel)
    finally:
        timer.join()
    assert not list(tmp_path.glob("*.xlsx"))
    assert not list(tmp_path.glob(".localocr-*"))


@pytest.mark.parametrize(
    "code", ["HTTP 401", "HTTP 403", "HTTP 404", "connection refused", "ENGINE_STARTUP_TIMEOUT"]
)
def test_batch_doctor_failure_reports_once_without_ocr(tmp_path, monkeypatch, code):
    provider = Provider(OpenWebUISettings(api="chat"))

    def failed():
        raise OCRServerError(code + " secret synthetic payload")

    monkeypatch.setattr(provider, "doctor", failed)
    monkeypatch.setattr(service, "create_provider", lambda *a, **kw: provider)
    folder = tmp_path / "inputs"
    folder.mkdir()
    for i in range(10):
        Image.new("RGB", (8, 8)).save(folder / f"{i}.png")
    out = tmp_path / "out"
    result = service.convert(folder, provider.settings, out, fail_soft=True)
    assert len(result) == 1 and result[0].batch_level and result[0].status == "FAILED"
    summary = json.loads((out / "summary.json").read_text())
    assert summary["counts"]["BLOCKED"] == summary["counts"]["INPUT_TOTAL"] == 10
    assert summary["counts"]["ATTEMPTED"] == summary["counts"]["XLSX_WITH_DATA"] == 0
    assert summary["batch"]["stage"] == "doctor"
    assert "secret synthetic payload" not in (out / "summary.json").read_text()
    assert not list(out.glob("*.xlsx"))


def test_exited_leader_descendants_and_unrelated_process(tmp_path):
    marker = tmp_path / "child.pid"
    code = 'import subprocess,sys; p=subprocess.Popen([sys.executable,"-c","import time; time.sleep(60)"]);open(sys.argv[1],"w").write(str(p.pid))'
    unrelated = launch_owned([sys.executable, "-c", "import time; time.sleep(60)"])
    process = launch_owned([sys.executable, "-c", code, str(marker)])
    try:
        deadline = time.monotonic() + 5
        while process.poll() is None and time.monotonic() < deadline:
            time.sleep(0.01)
        assert process.poll() == 0
        pid = int(marker.read_text())
        terminate_owned(process)
        assert unrelated.poll() is None
        if os.name == "posix":
            status = Path(f"/proc/{pid}/status")
            deadline = time.monotonic() + 2
            while (
                status.exists()
                and "\nState:\tZ" not in status.read_text()
                and time.monotonic() < deadline
            ):
                time.sleep(0.01)
            assert not status.exists() or "\nState:\tZ" in status.read_text()
        else:
            import ctypes
            from ctypes import wintypes

            api = ctypes.WinDLL("kernel32", use_last_error=True)
            api.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
            api.OpenProcess.restype = wintypes.HANDLE
            handle = api.OpenProcess(0x100000, False, pid)
            if handle:
                api.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
                api.CloseHandle.argtypes = [wintypes.HANDLE]
                assert api.WaitForSingleObject(handle, 2000) == 0
                api.CloseHandle(handle)
    finally:
        terminate_owned(process)
        terminate_owned(unrelated)


def test_writer_keeps_blank_rows_and_original_spaces(tmp_path):
    lines = ["", "  A  ", "", "   ", "B", ""]
    state = write_extraction(Extraction(lines=lines), tmp_path / "out.xlsx")
    book = load_workbook(tmp_path / "out.xlsx")
    sheet = book["Распознанный текст"]
    assert [(sheet.cell(i, 1).value or "") for i in range(1, 7)] == lines
    assert state["nonempty_data_cells"] == 2
    book.close()


def test_isolated_source_guard_cancel_and_no_overwrite(tmp_path):
    target = tmp_path / "out.xlsx"
    target.write_bytes(b"user workbook")
    with pytest.raises(FileExistsError):
        postprocess.convert_raw("наблюдение", target)
    assert target.read_bytes() == b"user workbook"
    cancel = Event()
    with pytest.raises(Cancelled):
        postprocess.convert_raw(
            "наблюдение", target, overwrite=True, cancel=cancel, before_publish=cancel.set
        )
    assert target.read_bytes() == b"user workbook"
    assert not list(tmp_path.glob(".localocr-*"))


def test_invalid_large_markdown_never_invents_coordinates():
    ex = extract("| A | B |\n| --- | --- |\n" + ("| Значение | " * 12000))
    assert not ex.document.tables
    assert ex.lines and "Значение" in "".join(ex.lines)


def test_stale_spool_cleanup_keeps_live_lease_and_foreign_paths(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    stale = work / "job-stale"
    live = work / "job-live"
    foreign = work / "job-foreign"
    staging = tmp_path / ".localocr-pending-stale"
    for path in [stale, live, foreign, staging]:
        path.mkdir()
    (stale / "raw.txt").write_text("synthetic private transient")
    (stale / "owner.json").write_text(
        json.dumps({"schema": "localocr-postprocess/1", "pending": str(staging)})
    )
    (staging / "owner.json").write_text(json.dumps({"spool": str(stale.resolve())}))
    (live / "owner.json").write_text(
        json.dumps({"schema": "localocr-postprocess/1", "pending": str(foreign)})
    )
    with postprocess.spool_lease(live):
        postprocess.recover_spools(work)
        assert live.exists() and foreign.exists()
        assert not stale.exists() and not staging.exists()


def test_own_job_close_is_idempotent(tmp_path):
    process = launch_owned([sys.executable, "-c", "pass"])
    terminate_owned(process)
    terminate_owned(process)
    assert process.poll() is not None


def test_actual_conversion_holds_recovery_lease(tmp_path):
    from localocr.runtime import portable_root

    work = portable_root() / "work/postprocess"

    def during_publication():
        jobs = list(work.glob("job-*"))
        assert len(jobs) == 1
        postprocess.recover_spools(work)
        assert jobs[0].exists()

    postprocess.convert_raw("observed", tmp_path / "out.xlsx", before_publish=during_publication)
    assert (tmp_path / "out.xlsx").is_file()
    assert not list(work.iterdir())


def test_term_resistant_owned_process_is_killed(tmp_path):
    from localocr.supervisor import run_owned

    code = "import signal,time;signal.signal(signal.SIGTERM,signal.SIG_IGN);time.sleep(60)"
    started = time.monotonic()
    with pytest.raises(TimeoutError):
        run_owned([sys.executable, "-c", code], cwd=tmp_path, timeout=0.3)
    assert time.monotonic() - started < 6
    assert run_owned([sys.executable, "-c", "pass"], cwd=tmp_path, timeout=3).exit_code == 0


@pytest.mark.parametrize(
    "detail,code",
    [
        ("Модель synthetic-model не установлена на сервере Ollama", "MODEL_UNAVAILABLE"),
        ("MODEL_NOT_FOUND/NOT_VISIBLE", "MODEL_UNAVAILABLE"),
        ("Embedded: CPU: runtime не установлен", "ENGINE_UNAVAILABLE"),
        ("Embedded: CUDA out of memory", "RESOURCE_FAILURE"),
    ],
)
def test_doctor_model_and_engine_have_distinct_safe_codes(detail, code):
    assert failsoft.failure_code(OCRServerError(detail), "doctor") == code
