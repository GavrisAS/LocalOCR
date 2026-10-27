"""Reliability contracts in the actual product service; no model or Internet required."""

import json
import os
import sys
from dataclasses import replace
from pathlib import Path
from threading import Event, Timer
from zipfile import BadZipFile

import pytest
from openpyxl import load_workbook
from PIL import Image
from PySide6.QtWidgets import QApplication
from typer.testing import CliRunner

from localocr.cli import app
from localocr.config import AppConfig, load_config, save_config
from localocr.domain import JobResult, StructureError
from localocr.failsoft import (
    extract,
    failure_code,
    selected_prompt,
    write_extraction,
)
from localocr.gui.app import SimpleWindow, Window
from localocr.ollama import Cancelled, OCRServerError, Settings
from localocr.providers import OpenWebUIChatProvider, OpenWebUISettings
from localocr.service import convert
from localocr.supervisor import run_owned

HTML = "<table><tr><td>Error</td><td>001</td></tr><tr><td></td><td>=1+1</td></tr></table>"
MD = "| Код | Цена |\n| --- | --- |\n| 001 | -1,20 |"


@pytest.mark.parametrize(
    "raw,structured,plain",
    [
        (HTML, True, False),
        (MD, True, False),
        ("до\n" + MD + "\nпосле", True, True),
        (HTML + "\n<table><tr><td>конец", True, True),
        ("<table><tr><td>A</td></tr><tr><td>B", False, True),
        ('<table><tr><td colspan="0">А</td></tr></table>', False, True),
        ('<table><tr><td rowspan="2">А</td></tr></table>', False, True),
        (HTML + HTML, True, False),
        ("  Счёт 001\n-0\n+4\n@A\n=1+1\n1,20\n01.02.2025", False, True),
        ("<think>private reasoning</think>Сумма 12", False, True),
        ("<|ref|>text<|/ref|><|det|>[[1,2,3,4]]<|/det|>Счёт", False, True),
    ],
)
def test_salvage_readback(tmp_path, raw, structured, plain):
    result = extract(raw)
    assert not result.outcome()["xlsx_with_data"]
    state = write_extraction(result, tmp_path / "out.xlsx")
    assert state["xlsx_with_data"] and state["status"] in {"PARTIAL", "REVIEW"}
    assert state["table_structured"] == structured
    assert state["partial_text_salvage"] == plain
    book = load_workbook(tmp_path / "out.xlsx")
    assert "Проверка" in book.sheetnames
    assert all(c.data_type != "f" for s in book for row in s for c in row)
    if raw == HTML:
        assert book["Таблица 1"]["A1"].value == "Error"
        assert book["Таблица 1"]["B1"].value == "001"
    book.close()


@pytest.mark.parametrize(
    "raw",
    [
        "",
        "   ",
        "<think>only reasoning",
        "<think>x</think>",
        "<script>secret()</script>",
        "<iframe>not OCR</iframe>",
        "<template>prompt</template>",
        "<|endoftext|>",
        "Table Recognition:",
        "Free OCR.",
        "<|grounding|>Convert the document to markdown.",
        "HTTP 401 Unauthorized",
        "Error: CUDA out of memory",
        "None None None",
        "№" * 80,
        "assistant",
    ],
)
def test_no_false_yield(tmp_path, raw):
    result = extract(raw)
    assert not result.values
    assert result.outcome()["error_codes"][-1] == "NO_CONTENT"
    with pytest.raises(ValueError, match="NO_CONTENT"):
        write_extraction(result, tmp_path / "out.xlsx")
    assert not (tmp_path / "out.xlsx").exists()


@pytest.mark.parametrize(
    "text",
    ["001", "+4", "-0", "-1,20", "1.200", "@X", "=SUM(A1:A2)", "000000000000000000123", "Error"],
)
def test_literal_values(tmp_path, text):
    result = extract("<table><tr><td>" + text + "</td></tr></table>")
    write_extraction(result, tmp_path / "out.xlsx")
    book = load_workbook(tmp_path / "out.xlsx")
    cell = book["Таблица 1"]["A1"]
    assert cell.value == text and cell.data_type == "s"
    book.close()


def test_xml_length_and_limits(tmp_path):
    text = "001\x00\\u0000\n" + "😀" * 20000
    ex = extract(text)
    write_extraction(ex, tmp_path / "out.xlsx")
    assert {"EXCEL_CELL_LIMIT_UNRESOLVED", "XML_ESCAPED"} <= set(ex.codes)
    book = load_workbook(tmp_path / "out.xlsx")
    lines = [r[0].value for r in book["Распознанный текст"]]
    assert lines[0] == "001\\u0000\\\\u0000"
    assert "".join(lines[1:]) == "😀" * 20000
    assert ex.outcome()["nonempty_data_cells"] == len(lines)
    book.close()
    with pytest.raises(ValueError, match="RESPONSE_LIMIT"):
        extract("a" * 4_000_001)


def test_atomic_corruption_overwrite_and_late_cancel(tmp_path, monkeypatch):
    import localocr.failsoft as f

    path = tmp_path / "out.xlsx"
    original = f.Workbook.save

    def corrupt(self, name):
        Path(name).write_bytes(b"broken")

    monkeypatch.setattr(f.Workbook, "save", corrupt)
    with pytest.raises(BadZipFile):
        write_extraction(extract("observed"), path)
    assert not list(tmp_path.iterdir())
    monkeypatch.setattr(f.Workbook, "save", original)
    path.write_bytes(b"previous")
    with pytest.raises(FileExistsError):
        write_extraction(extract("observed"), path)
    assert path.read_bytes() == b"previous"
    cancel = Event()

    def late(self, name):
        original(self, name)
        cancel.set()

    monkeypatch.setattr(f.Workbook, "save", late)
    with pytest.raises(Cancelled):
        write_extraction(extract("observed"), path, overwrite=True, cancel=cancel)
    assert path.read_bytes() == b"previous"
    assert list(tmp_path.iterdir()) == [path]


class Stub(OpenWebUIChatProvider):
    text = "наблюдение 001\n=1+1"
    calls = 0
    hook = None

    def doctor(self):
        self.model_identity = {"digest": "fixed"}
        return "ready"

    def request(self, method, endpoint, payload=None):
        self.calls += 1
        if self.hook:
            self.hook()
        self.http_ocr_calls = getattr(self, "http_ocr_calls", 0) + 1
        return {"choices": [{"message": {"content": self.text}, "finish_reason": "stop"}]}


def install_stub(monkeypatch, text=None):
    import localocr.service as s

    stub = Stub(OpenWebUISettings(api="chat"))
    if text is not None:
        stub.text = text
    monkeypatch.setattr(s, "create_provider", lambda *a, **kw: stub)
    return stub


def test_service_faults_cache_and_source_change(tmp_path, monkeypatch):
    stub = install_stub(monkeypatch)
    path = tmp_path / "image.png"
    Image.new("RGB", (8, 8), "white").save(path)
    settings = stub.settings
    out = tmp_path / "results"
    first = convert(path, settings, out, fail_soft=True)[0]
    assert first.xlsx_with_data and first.nonempty_data_cells == 2
    assert convert(path, settings, out, fail_soft=True)[0].resumed
    assert stub.calls == 1
    for changed in ["Free OCR.", "changed"]:
        assert not convert(path, settings, out, fail_soft=True, ocr_prompt=changed)[0].resumed
    with (out / "jobs.jsonl").open("a") as stream:
        stream.write('\n{"key":')
    assert convert(path, settings, out, fail_soft=True)[0].resumed
    stub.hook = lambda: path.write_bytes(b"changed")
    result = convert(path, settings, out / "changed", fail_soft=True)[0]
    assert result.status == "FAILED" and not result.output
    assert not list((out / "changed").glob("*.xlsx"))


def test_cancel_then_resume(tmp_path, monkeypatch):
    stub = install_stub(monkeypatch)
    path = tmp_path / "a.png"
    Image.new("RGB", (8, 8)).save(path)
    cancel = Event()
    stub.hook = cancel.set
    result = convert(path, stub.settings, tmp_path / "out", fail_soft=True, cancel=cancel)[0]
    assert result.status == "CANCELLED"
    assert result.calls == result.ocr_calls == result.http_attempts == 1
    assert not list((tmp_path / "out").glob("*.xlsx"))
    stub.hook = None
    assert convert(path, stub.settings, tmp_path / "out", fail_soft=True)[0].xlsx_with_data


@pytest.mark.parametrize(
    "error,stage,code",
    [
        (MemoryError(), "ocr", "RESOURCE_FAILURE"),
        (TimeoutError(), "ocr", "TIMEOUT"),
        (OCRServerError("READ_TIMEOUT"), "ocr", "TIMEOUT"),
        (OCRServerError("CONNECT_TIMEOUT"), "ocr", "TIMEOUT"),
        (OCRServerError("connection reset"), "ocr", "NETWORK_FAILURE"),
        *[
            (OCRServerError("HTTP " + str(n)), "ocr", c)
            for n, c in [
                (401, "HTTP_AUTH_401"),
                (403, "HTTP_FORBIDDEN_403"),
                (404, "HTTP_NOT_FOUND_404"),
                (429, "HTTP_RATE_LIMIT_429"),
                (500, "HTTP_SERVER_5XX"),
                (503, "HTTP_SERVER_5XX"),
            ]
        ],
        (OSError(), "xlsx", "IO_FAILURE"),
        (ValueError(), "load", "IMAGE_DECODE_FAILURE"),
        (StructureError(), "ocr", "PARSER_FAILURE"),
    ],
)
def test_fixed_error_codes(error, stage, code):
    assert failure_code(error, stage) == code


def test_supervisor_crash_timeout_cancel_then_next(tmp_path):
    result = run_owned([sys.executable, "-c", "import os; os._exit(7)"], cwd=tmp_path, timeout=2)
    assert result.worker_crashed and result.exit_code == 7
    with pytest.raises(TimeoutError):
        run_owned([sys.executable, "-c", "import time; time.sleep(10)"], cwd=tmp_path, timeout=0.1)
    cancel = Event()
    timer = Timer(0.1, cancel.set)
    timer.start()
    try:
        with pytest.raises(Cancelled):
            run_owned(
                [sys.executable, "-c", "import time; time.sleep(10)"],
                cwd=tmp_path,
                timeout=2,
                cancel=cancel,
            )
    finally:
        timer.join()
    result = run_owned([sys.executable, "-c", 'print("private-output")'], cwd=tmp_path, timeout=2)
    assert result.exit_code == 0 and "private-output" not in repr(result)


@pytest.mark.skipif(os.name == "nt", reason="POSIX process-group assertion; Windows kit required")
def test_supervisor_reaps_descendant(tmp_path):
    code = 'import subprocess,sys,time; p=subprocess.Popen([sys.executable,"-c","import time; time.sleep(20)"]);open("child.pid","w").write(str(p.pid));time.sleep(20)'
    with pytest.raises(TimeoutError):
        run_owned([sys.executable, "-c", code], cwd=tmp_path, timeout=0.3)
    pid = int((tmp_path / "child.pid").read_text())
    status = Path(f"/proc/{pid}/status")
    assert not status.exists() or "\nState:\tZ" in status.read_text()


@pytest.mark.parametrize("window_type", [Window, SimpleWindow])
def test_gui_opt_in_and_preserves_chat(window_type):
    appqt = QApplication.instance() or QApplication([])
    save_config(
        replace(
            AppConfig(),
            openwebui_api="chat",
            openwebui_url="http://127.0.0.1:1",
            openwebui_model="deepseek-ocr",
            api_token="fixture-token",
        )
    )
    window = window_type()
    assert not window.fail_soft.isChecked()
    window.fail_soft.setChecked(True)
    assert window.current_config().fail_soft
    assert window.runtime_settings().api == "chat"
    window.hide()
    window.deleteLater()
    appqt.processEvents()


def test_cli_opt_in_and_old_schema(tmp_path, monkeypatch):
    assert not load_config().fail_soft
    assert not JobResult("old", "PARTIAL").xlsx_with_data
    calls = []

    def run(*a, **kw):
        calls.append(kw["fail_soft"])
        return []

    monkeypatch.setattr("localocr.cli.run_convert", run)
    for flags, expected in [([], False), (["--fail-soft"], True), (["--no-fail-soft"], False)]:
        result = CliRunner().invoke(app, ["convert", str(tmp_path), *flags])
        assert result.exit_code == 0, result.output
        assert calls[-1] == expected
    assert selected_prompt(Settings(model="deepseek-ocr:3b")).startswith("<|grounding|>")
    assert selected_prompt(Settings()) == "Table Recognition:"


def test_final_source_guard_and_merge_readback(tmp_path):
    path = tmp_path / "out.xlsx"
    ex = extract('<table><tr><td colspan="2">001</td></tr></table>')

    def changed():
        raise ValueError("SOURCE_CHANGED")

    with pytest.raises(ValueError, match="SOURCE_CHANGED"):
        write_extraction(ex, path, before_publish=changed)
    assert not path.exists()
    write_extraction(ex, path)
    book = load_workbook(path)
    assert [str(r) for r in book["Таблица 1"].merged_cells.ranges] == ["A1:B1"]
    book.close()


@pytest.mark.parametrize(
    "raw",
    [
        "assistant " * 40,
        "  Table Recognition:  ",
        "<analysis>reasoning</analysis>",
        "<h1>500 Internal Server Error</h1>",
        "generation time: 1.23s",
        "\x00\x01\u200b",
        "<|ref|>table<|/ref|><|det|>[[1,2,3,4]]<|/det|>",
    ],
)
def test_more_artifact_envelopes(raw):
    assert not extract(raw).values


def test_large_structure_retains_text_and_self_closing_html():
    ex = extract('<table><tr><td rowspan="10000" colspan="1000">001</td></tr></table>')
    assert not ex.document.tables and ex.lines == ["001"]
    assert ex.codes == ["RESPONSE_STRUCTURE_LIMIT"]
    assert extract("<iframe/>\nполезный текст").lines == ["полезный текст"]
    assert extract("<think/>\nполезный текст").lines == ["полезный текст"]
    with pytest.raises(ValueError, match="RESPONSE_LINE_LIMIT"):
        extract("x\n" * 100002)


@pytest.mark.parametrize("status", [200, 500, 503])
def test_server_memory_error_is_private_and_not_retried(server, status):
    from localocr.ollama import OllamaClient
    from localocr.providers import create_provider

    lines = []
    server[1]["status"] = status
    server[1]["body"] = {"error": "CUDA out of memory; private upstream fixture"}
    client = create_provider(Settings(url=server[0], retries=1), on_network_event=lines.append)
    assert isinstance(client, OllamaClient)
    try:
        with pytest.raises(OCRServerError) as error:
            client.request(
                "POST", "/api/generate", {"images": ["fixture"], "model": "glm-ocr:latest"}
            )
        assert failure_code(error.value, "ocr") == "RESOURCE_FAILURE"
        assert "private upstream fixture" not in str(error.value) + "\n".join(lines)
        assert len(server[1]["requests"]) == 1
    finally:
        client.close()


def test_deleted_queued_file_has_terminal_record_and_next_file(tmp_path, monkeypatch):
    stub = install_stub(monkeypatch)
    folder = tmp_path / "inputs"
    folder.mkdir()
    for name in ["a.png", "b.png"]:
        Image.new("RGB", (8, 8)).save(folder / name)

    def progress(done, total, message):
        if done == 0:
            (folder / "a.png").unlink(missing_ok=True)

    out = tmp_path / "results"
    rows = convert(folder, stub.settings, out, fail_soft=True, progress=progress)
    assert rows[0].status == "FAILED" and rows[1].xlsx_with_data
    records = [
        json.loads(line) for line in (out / "jobs.jsonl").read_text().splitlines() if line.strip()
    ]
    assert any(r.get("source_unavailable") and r["status"] == "FAILED" for r in records)


def test_error_inside_invalid_html_still_observed():
    ex = extract('<table><tr><td colspan="0"><b>Error</b></td><td>001</td></tr></table>')
    assert ex.lines == ["Error", "001"] and not ex.document.tables


def test_cancel_during_final_source_check_never_publishes(tmp_path):
    cancel = Event()
    with pytest.raises(Cancelled):
        write_extraction(
            extract("observed"), tmp_path / "out.xlsx", cancel=cancel, before_publish=cancel.set
        )
    assert not list(tmp_path.iterdir())


def test_published_hash_io_failure_does_not_abort_next_file(tmp_path, monkeypatch):
    import localocr.service as service

    stub = install_stub(monkeypatch)
    folder = tmp_path / "inputs"
    folder.mkdir()
    for name in ["a.png", "b.png"]:
        Image.new("RGB", (8, 8)).save(folder / name)
    original_hash = service.file_hash

    def hashing(path):
        if path.name == "a.xlsx":
            raise PermissionError("bounded fixture")
        return original_hash(path)

    monkeypatch.setattr(service, "file_hash", hashing)
    rows = convert(folder, stub.settings, tmp_path / "results", fail_soft=True)
    assert rows[0].status == "FAILED" and rows[0].error_codes == ["IO_FAILURE"]
    assert rows[1].xlsx_with_data
