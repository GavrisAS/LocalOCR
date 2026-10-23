"""Independent content/position, partial publication and durable batch invariants."""

import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from localocr.domain import Cell, CellProvenance, StructureError, Table, TableDocument
from localocr.excel import write_workbook
from localocr.imaging import load_image, source_polygon
from localocr.journal import Journal, atomic_json
from localocr.ollama import Cancelled, Settings
from localocr.reconstruction import recognize_document
from localocr.service import convert
from localocr.synthetic import generate


def ruled(rows=3, columns=3, texts=None):
    image = Image.new("RGB", (columns * 100 + 21, rows * 60 + 21), "white")
    draw = ImageDraw.Draw(image)
    for x in range(columns + 1):
        draw.line((10 + x * 100, 10, 10 + x * 100, 10 + rows * 60), fill="black", width=2)
    for y in range(rows + 1):
        draw.line((10, 10 + y * 60, 10 + columns * 100, 10 + y * 60), fill="black", width=2)
    for row, column, text in texts or []:
        draw.text((20 + column * 100, 30 + row * 60), text, fill="black")
    return image


def html(rows):
    return (
        "<table>"
        + "".join("<tr>" + "".join(f"<td>{text}</td>" for text in row) + "</tr>" for row in rows)
        + "</table>"
    )


class Client:
    settings = SimpleNamespace(max_calls=12, max_image_seconds=30)

    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = 0

    def check_cancel(self):
        pass

    def recognize(self, image, *, single_table=False):
        self.calls += 1
        result = next(self.responses)
        if isinstance(result, Exception):
            raise result
        return result


def test_omitted_nonempty_row_is_recovered_without_shifting_or_text_deduplication():
    # Identical row contents are legitimate distinct source rows.
    image = ruled(3, 2, [(r, 0, "same") for r in range(3)])
    client = Client(
        [
            html([["same", "1"], ["same", "3"]]),
            html([["same", "1"]]),
            html([["same", "2"], ["same", "3"]]),
        ]
    )
    doc = recognize_document(image, client, None)
    assert [
        (doc.tables[0].grid()[r, 0].row, doc.tables[0].grid()[r, 1].text) for r in range(3)
    ] == [(0, "1"), (1, "2"), (2, "3")]
    assert doc.outcome()["status"] == "REVIEW"
    assert doc.outcome()["verified_anchors"] == 0
    assert client.calls == 3


def test_missing_td_does_not_shift_neighbor_into_blank_position():
    client = Client(
        [html([["left", "right"], ["a", "b", "c"], ["d", "e", "f"]]), html([["left", "", "right"]])]
    )
    doc = recognize_document(ruled(), client, None)
    grid = doc.tables[0].grid()
    assert (grid[0, 0].text, grid[0, 1].text, grid[0, 2].text) == ("left", "", "right")
    assert [grid[1, c].text for c in range(3)] == ["a", "b", "c"]
    assert grid[0, 1].provenance.state == "observed-empty"


def test_visible_unrecognized_ink_is_partial_and_xlsx_has_no_fake_values(tmp_path):
    image = ruled(2, 2, [(0, 1, "visible"), (1, 0, "visible")])
    client = Client(
        [html([["=1+1", ""], ["", "12,50"]]), html([["=1+1", ""]]), html([["", "12,50"]])]
    )
    doc = recognize_document(image, client, None)
    outcome = doc.outcome()
    assert outcome["status"] == "PARTIAL"
    assert outcome["recognized_anchors"] == 2
    assert outcome["unresolved_regions"] == 2
    target = tmp_path / "partial.xlsx"
    write_workbook(doc, target)
    book = load_workbook(target)
    sheet = book["Таблица 1"]
    assert (sheet["A1"].value, sheet["A1"].data_type) == ("=1+1", "s")
    assert sheet["B1"].value is None and sheet["A2"].value is None
    assert sheet["B2"].value == 12.5 and sheet["B2"].number_format == "0.00"
    assert "Проверка" in book and book["Проверка"]["E2"].value
    assert all(cell.comment is None for sheet in book for row in sheet for cell in row)
    book.close()


def test_invalid_span_is_local_and_other_rows_survive():
    response = "<table><tr><td rowspan='99'>bad</td><td>h</td></tr><tr><td>a</td><td>b</td></tr><tr><td>c</td><td>d</td></tr></table>"
    client = Client([response, html([["h", "j"]]), html([["a", "b"], ["c", "d"]])])
    doc = recognize_document(ruled(3, 2), client, None)
    assert [doc.tables[0].grid()[r, 1].text for r in range(3)] == ["j", "b", "d"]
    assert not any(cell.rowspan > 3 for cell in doc.tables[0].cells)


def test_total_request_budget_and_no_empty_workbook(tmp_path):
    client = Client(["not a table"] * 20)
    client.settings = SimpleNamespace(max_calls=2, max_image_seconds=30)
    with pytest.raises(StructureError):
        recognize_document(ruled(), client, tmp_path)
    assert client.calls == 2
    assert json.loads((tmp_path / "run.json").read_text())["status"] == "FAILED"
    assert not list(tmp_path.glob("*.xlsx"))


def test_cancel_keeps_existing_output_and_removes_temporary_workbook(tmp_path):
    from threading import Event

    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"previous output")
    cancel = Event()
    cancel.set()
    with pytest.raises(Cancelled):
        write_workbook(TableDocument([Table([Cell(0, 0, "new")], 1, 1)]), target, True, cancel)
    assert target.read_bytes() == b"previous output"
    assert not list(tmp_path.glob(".localocr-*.xlsx"))


def test_journal_recovers_torn_append_and_excludes_concurrent_writer(tmp_path):
    with Journal(tmp_path) as journal:
        journal.append({"key": "done", "status": "SUCCESS"})
        with pytest.raises(ValueError, match="занят"):
            with Journal(tmp_path):
                pass
    with (tmp_path / "jobs.jsonl").open("ab") as stream:
        stream.write(b'{"key": "torn')
    with Journal(tmp_path) as journal:
        assert journal.records["done"]["status"] == "SUCCESS"
        journal.append({"key": "after-crash", "status": "RUNNING"})
    with Journal(tmp_path) as journal:
        assert "after-crash" in journal.records and "torn" not in journal.records
    atomic_json(tmp_path / "report.json", {"completed": True})
    assert json.loads((tmp_path / "report.json").read_text()) == {"completed": True}


def test_batch_resume_invalidation_and_tampered_output(server, tmp_path):
    from unittest.mock import patch

    source = tmp_path / "image.png"
    ruled(2, 2).save(source)
    # This transport fixture represents a declared digest, not real inference.
    original = __import__("localocr.ollama", fromlist=["OllamaClient"]).OllamaClient.doctor

    def doctor(self):
        result = original(self)
        self.model_identity["digest"] = "synthetic-digest"
        return result

    settings = Settings(url=server[0])
    with patch("localocr.ollama.OllamaClient.doctor", doctor):
        first = convert(source, settings)[0]
        calls = sum(payload is not None for _, payload in server[1]["requests"])
        second = convert(source, settings)[0]
        assert second.resumed and second.status == "SKIPPED"
        assert sum(payload is not None for _, payload in server[1]["requests"]) == calls
        changed = convert(source, replace(settings, num_predict=8000))[0]
        assert changed.output != first.output and changed.status == "REVIEW"
        assert first.output is not None and changed.output is not None
        assert Path(first.output).exists()
        Path(changed.output).write_bytes(b"tampered")
        check = convert(source, replace(settings, num_predict=8000))[0]
        assert not check.resumed


@pytest.mark.parametrize("orientation", range(1, 9))
def test_source_coordinates_survive_exif(tmp_path, orientation):
    image = Image.new("RGB", (100, 60), "white")
    exif = Image.Exif()
    exif[274] = orientation
    path = tmp_path / "source.png"
    image.save(path, exif=exif)
    result = load_image(path)
    corners = source_polygon(result, (0, 0, result.width, result.height))
    assert set(corners) == {(0, 0), (100, 0), (100, 60), (0, 60)}


def test_seeded_holdout_has_valid_random_merges_and_varied_styles():
    cases = [generate(seed) for seed in range(40)]
    assert {c.style for c in cases} == {"grid", "borderless", "open", "broken"}
    assert any(c.rotation == 90 for c in cases)
    assert any(cell.rowspan > 1 for c in cases for cell in c.gold.cells)
    assert any(cell.colspan > 1 for c in cases for cell in c.gold.cells)
    assert any(not cell.text for c in cases for cell in c.gold.cells)
    for case in cases:
        assert len(case.gold.grid()) == case.gold.rows * case.gold.columns
    first, second = generate(61), generate(61)
    assert first.image.tobytes() == second.image.tobytes()
    assert first.gold == second.gold


def test_blank_unresolved_cell_can_never_count_as_recognized_data():
    cell = Cell(0, 0, "", provenance=CellProvenance("unresolved"))
    assert TableDocument([Table([cell], 1, 1)]).outcome()["status"] == "FAILED"


def test_transport_retries_respect_total_physical_ocr_budget(server):
    from localocr.ollama import OCRServerError, OllamaClient

    url, state = server
    state["status"] = 500
    client = OllamaClient(Settings(url=url, retries=1))
    client.request_budget_remaining = 1
    with pytest.raises(OCRServerError, match="BUDGET"):
        client.recognize(b"synthetic")
    assert client.http_ocr_calls == 1
    assert len(state["requests"]) == 1


def test_ollama_expert_options_and_usage_are_observed(server):
    from localocr.ollama import OllamaClient

    url, state = server
    state["body"] = {
        "response": "<table><tr><td>A</td></tr></table>",
        "done": True,
        "done_reason": "stop",
        "eval_count": 11,
        "prompt_eval_count": 35,
    }
    client = OllamaClient(Settings(url=url, num_ctx=65536, num_predict=8000, table_stop=False))
    client.recognize(b"synthetic", single_table=True)
    options = state["requests"][-1][1]["options"]
    assert options["num_ctx"] == 65536 and options["num_predict"] == 8000
    assert "stop" not in options
    assert client.last_diagnostic["eval_count"] == 11
    assert client.last_diagnostic["prompt_eval_count"] == 35
    assert client.last_diagnostic["thinking_output"] is False


def test_unavailable_backend_is_not_retried_for_every_crop():
    from localocr.ollama import OCRServerError

    client = Client([OCRServerError("CONNECT_REFUSED")] * 20)
    with pytest.raises(StructureError):
        recognize_document(ruled(8, 3), client, None)
    assert client.calls == 1
