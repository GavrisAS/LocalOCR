from threading import Event
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment

from localocr.domain import Cell, StructureError, Table, TableDocument, ValidationIssue
from localocr.excel import _autofit_sheet, value, write_workbook
from localocr.ollama import Cancelled
from localocr.quality import compare, export
from localocr.tables import parse_response


@pytest.mark.parametrize(
    "text,expected",
    [
        ("0012", "0012"),
        ("5,50", 5.5),
        ("2590-254", "2590-254"),
        ("0", 0),
        ("-4", -4),
        ("=1+1", "=1+1"),
        ("1234567890123456", "1234567890123456"),
        ("1e4", "1e4"),
        ("+123", "+123"),
        ("1.20", 1.2),
        ("01.20", "01.20"),
        ("∞", "∞"),
    ],
)
def test_semantics(text, expected):
    assert value(text) == expected


def test_roundtrip(tmp_path):
    doc = parse_response(
        '<table><tr><th rowspan="2">Код</th><th colspan="2">Заголовок</th></tr><tr><th>А</th><th>Б</th></tr><tr><td>0012</td><td>=1+1</td><td></td></tr></table>'
    )
    doc.tables[0].issues.append(ValidationIssue("TEST", "Проверить"))
    path = tmp_path / "выход.xlsx"
    write_workbook(doc, path)
    book = load_workbook(path)
    sheet = book.worksheets[0]
    assert (sheet.max_row, sheet.max_column) == (3, 3)
    assert {str(r) for r in sheet.merged_cells.ranges} == {"A1:A2", "B1:C1"}
    assert sheet["A3"].value == "0012"
    assert sheet["B3"].value == "=1+1" and sheet["B3"].data_type == "s"
    assert sheet["C3"].value is None
    assert sheet["A1"].font.bold and sheet["A1"].alignment.wrap_text
    assert sheet["A3"].border.left.style == "thin"
    assert all(cell.comment is None for row in sheet for cell in row)
    assert sheet.freeze_panes == "A3"
    assert "Проверка" in book.sheetnames
    book.close()
    with ZipFile(path) as archive:
        assert not any(
            name.startswith("xl/comments/")
            or name.startswith("xl/drawings/vmlDrawing")
            or (name.endswith(".rels") and b"/comments" in archive.read(name))
            for name in archive.namelist()
        )
    original = path.read_bytes()
    with pytest.raises(FileExistsError):
        write_workbook(doc, path)
    assert path.read_bytes() == original
    event = Event()
    event.set()
    with pytest.raises(Cancelled):
        write_workbook(doc, tmp_path / "cancel.xlsx", cancel=event)
    assert not (tmp_path / "cancel.xlsx").exists()
    assert not list(tmp_path.glob(".localocr-*"))


@pytest.mark.parametrize("rows,columns", [(2, 2), (3, 5), (9, 32), (12, 50)])
def test_dimensions(tmp_path, rows, columns):
    doc = TableDocument(
        [
            Table(
                [Cell(r, c, f"{r}:{c}") for r in range(rows) for c in range(columns)], rows, columns
            )
        ]
    )
    path = tmp_path / "result.xlsx"
    write_workbook(doc, path)
    book = load_workbook(path)
    assert (book.active.max_row, book.active.max_column) == (rows, columns)
    assert book.active.cell(rows, columns).value == f"{rows - 1}:{columns - 1}"
    book.close()


def test_invalid_not_published(tmp_path):
    with pytest.raises(StructureError):
        write_workbook(TableDocument([Table([Cell(0, 0, "x")], 2, 2)]), tmp_path / "bad.xlsx")
    assert not list(tmp_path.iterdir())


def test_review_text_cannot_be_formula(tmp_path):
    doc = TableDocument(
        [
            Table(
                [Cell(0, 0, "=1+1", warnings=["=2+2"])],
                1,
                1,
                issues=[ValidationIssue("x", '=HYPERLINK("bad")')],
            )
        ]
    )
    path = tmp_path / "safe.xlsx"
    write_workbook(doc, path)
    book = load_workbook(path)
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    assert book["Проверка"]["D2"].value == "=2+2"
    book.close()


@pytest.mark.parametrize(
    "text",
    ["5,50", "5,70", "9,60", "7,00", "24,00", "0,5", "0012", "2590-254", "1e4", "01.02", "=1+1"],
)
def test_printed_fidelity(tmp_path, text):
    path = tmp_path / "fidelity.xlsx"
    write_workbook(TableDocument([Table([Cell(0, 0, text)], 1, 1)]), path)
    book = load_workbook(path)
    cell = book.active["A1"]
    assert cell.comment is None
    if "," in text:
        places = len(text.split(",")[1])
        assert cell.data_type == "n"
        assert cell.number_format == "0." + "0" * places
        assert f"{cell.value:.{places}f}".replace(".", ",") == text
    else:
        assert cell.value == text and cell.data_type == "s"
    book.close()


def test_content_autofit_and_merged_dimensions(tmp_path):
    doc = TableDocument(
        [
            Table(
                [
                    Cell(0, 0, "Наименование", is_header=True),
                    Cell(0, 1, "Цена", is_header=True),
                    Cell(1, 0, "Очень длинное название на кириллице с несколькими словами"),
                    Cell(1, 1, "5,50"),
                    Cell(2, 0, "Первая строка\nВторая строка"),
                    Cell(2, 1, "0012345678901234567890123456789012345678901234567890"),
                    Cell(3, 0, "Общий многоуровневый заголовок", rowspan=1, colspan=2),
                ],
                4,
                2,
            )
        ]
    )
    path = tmp_path / "autofit.xlsx"
    write_workbook(doc, path)
    book = load_workbook(path)
    sheet = book.active
    assert {str(area) for area in sheet.merged_cells.ranges} == {"A4:B4"}
    assert 8 <= sheet.column_dimensions["A"].width <= 52
    assert 8 <= sheet.column_dimensions["B"].width <= 52
    assert sheet.row_dimensions[2].height > 15
    assert sheet.row_dimensions[3].height > 15
    assert sheet.row_dimensions[4].height > 15
    assert sheet["B2"].value == 5.5 and sheet["B2"].number_format == "0.00"
    assert sheet["B3"].value == "0012345678901234567890123456789012345678901234567890"
    book.close()


def test_dimensions_ignore_image_geometry(tmp_path):
    docs = [
        TableDocument(
            [
                Table(
                    [Cell(0, 0, "Заголовок"), Cell(1, 0, "длинный текст с переносом")],
                    2,
                    1,
                    row_heights=heights,
                    column_widths=widths,
                )
            ]
        )
        for heights, widths in [([12.0, 20.0], [8.0]), ([180.0, 180.0], [500.0])]
    ]
    dimensions = []
    for index, doc in enumerate(docs):
        path = tmp_path / f"geometry-{index}.xlsx"
        write_workbook(doc, path)
        book = load_workbook(path)
        dimensions.append(
            (book.active.column_dimensions["A"].width, book.active.row_dimensions[2].height)
        )
        book.close()
    assert dimensions[0] == dimensions[1]


def test_rotated_header_uses_vertical_height():
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Вертикальный заголовок"
    sheet["A1"].alignment = Alignment(textRotation=255, wrap_text=True)
    _autofit_sheet(sheet)
    assert sheet.column_dimensions["A"].width < 20
    assert sheet.row_dimensions[1].height > 100
    book.close()


def test_display_limit_is_reported_and_currency_is_measured(tmp_path):
    path = tmp_path / "display-limit.xlsx"
    write_workbook(TableDocument([Table([Cell(0, 0, "x" * 4000)], 1, 1)]), path)
    book = load_workbook(path)
    assert book.active.column_dimensions["A"].width == 52
    assert book.active.row_dimensions[1].height == 150
    assert "DISPLAY_LIMIT" in book["Проверка"]["D2"].value
    book.close()

    currency_widths = []
    for number_format in ["#,##0.00", '"₽"#,##0.00']:
        currency_book = Workbook()
        currency_sheet = currency_book.active
        currency_sheet["A1"] = 1234.5
        currency_sheet["A1"].number_format = number_format
        _autofit_sheet(currency_sheet)
        currency_widths.append(currency_sheet.column_dimensions["A"].width)
        currency_book.close()
    assert currency_widths[1] > currency_widths[0]


def test_quality_sidecar_and_legacy_comments(tmp_path):
    path = tmp_path / "quality.xlsx"
    write_workbook(TableDocument([Table([Cell(0, 0, "5,50")], 1, 1)]), path)
    actual = export(path)
    assert actual["tables"][0]["cells"][0]["text"] is None
    gold = {
        "human_verified": True,
        "tables": [
            {
                "rows": 1,
                "columns": 1,
                "cells": [
                    {
                        "row": 1,
                        "column": 1,
                        "spans": [1, 1],
                        "text": "5,50",
                        "type": "n",
                    }
                ],
            }
        ],
    }
    result = compare(actual, gold)
    assert result["status"] == "RAW_TEXT_UNAVAILABLE"
    assert result["counts"]["original_text_unavailable"] == 1
    assert result["counts"].get("text_mismatches", 0) == 0

    diagnostics = tmp_path / "diagnostics"
    diagnostics.mkdir()
    (diagnostics / "observations.json").write_text(
        '{"schema_version":1,"tables":[{"rows":1,"columns":1,"cells":'
        '[{"row":0,"column":0,"rowspan":1,"colspan":1,"text":"5,50"}]}]}',
        encoding="utf-8",
    )
    assert export(path, diagnostics)["tables"][0]["cells"][0]["text"] == "5,50"

    historical = load_workbook(path)
    historical.active["A1"].comment = Comment("5,50", "LocalOCR original OCR text")
    historical.save(path)
    historical.close()
    assert export(path)["tables"][0]["cells"][0]["text"] == "5,50"


def test_hardlink_failure_and_race(tmp_path, monkeypatch):
    import os

    doc = TableDocument([Table([Cell(0, 0, "x")], 1, 1)])
    path = tmp_path / "result.xlsx"

    def unsupported(*args):
        raise OSError("unsupported")

    monkeypatch.setattr(os, "link", unsupported)
    with pytest.raises(OSError, match="hardlinks unsupported"):
        write_workbook(doc, path)
    assert not path.exists() and not list(tmp_path.glob(".localocr-*"))

    def race(*args):
        path.write_bytes(b"competitor")
        raise FileExistsError()

    monkeypatch.setattr(os, "link", race)
    with pytest.raises(FileExistsError):
        write_workbook(doc, path)
    assert path.read_bytes() == b"competitor"
    assert not list(tmp_path.glob(".localocr-*"))


def test_overwrite_cancel_and_replace_failure(tmp_path, monkeypatch):
    import os

    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original")
    doc = TableDocument([Table([Cell(0, 0, "x")], 1, 1)])
    cancel = Event()
    cancel.set()
    with pytest.raises(Cancelled):
        write_workbook(doc, path, overwrite=True, cancel=cancel)
    assert path.read_bytes() == b"original"

    def denied(*args):
        raise PermissionError("denied")

    monkeypatch.setattr(os, "replace", denied)
    with pytest.raises(PermissionError):
        write_workbook(doc, path, overwrite=True)
    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob(".localocr-*"))
