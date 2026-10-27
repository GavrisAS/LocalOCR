"""Candidate integration through the shared product service and transport interface."""

import json
from threading import Event

import pytest
from openpyxl import load_workbook
from PIL import Image

from localocr.domain import TableDocument
from localocr.failsoft import FailSoftDocument, recognize_failsoft
from localocr.ollama import Cancelled
from localocr.providers import OpenWebUIChatProvider, OpenWebUISettings
from localocr.service import convert


class Stub(OpenWebUIChatProvider):
    def doctor(self):
        self.model_identity = {"digest": "fixed-test-digest"}
        return "ready"

    def request(self, method, endpoint, payload=None):
        assert endpoint == "/api/chat/completions"
        assert isinstance(payload, dict)
        assert payload["messages"][0]["content"][0]["text"] == "Free OCR."
        assert payload["messages"][0]["content"][1]["image_url"]["url"].startswith(
            "data:image/png;base64,"
        )
        return {
            "choices": [{"message": {"content": "001\n=1+1\n-0\n1,20"}, "finish_reason": "length"}]
        }


def test_candidate_same_service_resume_and_truncated_raw(tmp_path, monkeypatch):
    import localocr.service as service

    settings = OpenWebUISettings(url="http://127.0.0.1:1", api="chat", num_predict=7800)
    monkeypatch.setattr(service, "create_provider", lambda settings, cancel: Stub(settings, cancel))
    image = tmp_path / "input.png"
    Image.new("RGB", (16, 16), "white").save(image)
    out = tmp_path / "results"
    first = convert(
        image,
        settings,
        out,
        diagnostics=tmp_path / "private",
        fail_soft=True,
        ocr_prompt="Free OCR.",
    )[0]
    assert first.status == "PARTIAL"
    assert first.recognized_anchors == 0
    assert first.partial_text_salvage and first.xlsx_with_data
    assert "SUSPECTED_TRUNCATION" in first.error_codes
    book = load_workbook(first.output)
    assert book.active.title == "Проверка"
    assert [r[0].value for r in book["Распознанный текст"]] == ["001", "=1+1", "-0", "1,20"]
    assert all(c.data_type != "f" for s in book for row in s for c in row)
    book.close()
    second = convert(image, settings, out, fail_soft=True, ocr_prompt="Free OCR.")[0]
    assert second.resumed and second.status == "PARTIAL"
    assert second.error_codes == first.error_codes
    assert (out / "summary.json").exists()
    assert (out / "summary.csv").exists()
    assert json.loads(next(out.glob("report-*.json")).read_text())["results"]
    assert (tmp_path / "private/input/response-1.txt").exists()


def test_dto_defaults_and_cancel(tmp_path):
    assert TableDocument([]).outcome()["status"] == "FAILED"
    doc = FailSoftDocument([], unplaced_lines=["unplaced"])
    doc.validate()
    assert doc.outcome()["status"] == "PARTIAL"
    assert doc.outcome()["recognized_anchors"] == 0
    cancel = Event()
    cancel.set()
    client = Stub(OpenWebUISettings(api="chat"), cancel)
    with pytest.raises(Cancelled):
        recognize_failsoft(Image.new("RGB", (8, 8)), client, None, prompt="Free OCR.")
