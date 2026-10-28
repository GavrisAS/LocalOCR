"""Stage210 contracts use synthetic documents; no private data or model inference."""

import hashlib
import io
import json
import runpy
from dataclasses import asdict, replace
from pathlib import Path
from threading import Event, Timer

import pytest
from openpyxl import load_workbook
from PIL import Image

from localocr import deepseek_layout, failsoft, postprocess, service
from localocr.domain import JobResult, StructureError
from localocr.failsoft import extract, failure_code, recognize_failsoft, write_extraction
from localocr.journal import Journal, settings_identity
from localocr.ollama import Cancelled, OCRServerError, Settings
from localocr.providers import (
    OpenWebUIChatProvider,
    OpenWebUIOllamaProxyProvider,
    OpenWebUISettings,
)
from localocr.response_protocol import completion_metadata
from localocr.tables import parse_response

FIXTURES = json.loads(Path("tests/fixtures/stage210.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("case", FIXTURES, ids=lambda c: c["id"])
def test_fixture_contract_and_determinism(case, tmp_path):
    result = extract(case["raw"], model_id=case["model"])
    assert result.values == case["values"]
    assert len(result.document.tables) == case["tables"]
    assert len(result.layout_metadata) == case["removed"]
    assert result.outcome()["layout_markers_remaining"] == (1 if case["id"] == "DS-05" else 0)
    assert asdict(result) == asdict(extract(case["raw"], model_id=case["model"]))
    if not result.values:
        assert result.outcome()["no_content"]
        with pytest.raises(ValueError, match="NO_CONTENT"):
            write_extraction(result, tmp_path / "result.xlsx")
        return
    state = write_extraction(result, tmp_path / "result.xlsx")
    assert state["ocr_contentful"] and state["xlsx_with_data"]
    book = load_workbook(tmp_path / "result.xlsx")
    values = [c.value for s in book if s.title != "Проверка" for row in s for c in row if c.value]
    assert sorted(values) == sorted(case["values"])
    assert all(c.data_type != "f" for s in book for row in s for c in row)
    book.close()


@pytest.mark.parametrize("label", ["table", "text", "sub_title"])
@pytest.mark.parametrize("space", ["", " ", "\t", "\r\n"])
def test_layout_property_like_metadata_payload_and_idempotency(label, space):
    for start in range(0, 1000, 71):
        raw = f"{label}{space}[[ {start},0,{start + 1},1000 ]]0012\nРусский =1+1"
        adapted = deepseek_layout.adapt(raw)
        assert len(adapted.metadata) == 1
        assert deepseek_layout.adapt(adapted.text).text == adapted.text
        assert extract(raw, model_id="deepseek-ocr:3b").values == ["0012", "Русский =1+1"]


@pytest.mark.parametrize(
    "bbox",
    [
        "1,2,3",
        "a,2,3,4",
        "-1,2,3,4",
        "1,2,1001,4",
        "3,2,1,4",
        "1,2,1,4",
        "[1,2,3,4]",
        "1,2,[3,4",
        "1,2,3,4",
    ],
)
def test_uncertain_layout_is_not_silently_erased(bbox):
    raw = "text[[" + bbox  # missing closure, nested syntax and invalid fields remain literal
    result = extract(raw, model_id="deepseek-ocr:3b")
    assert result.values == [raw]
    assert "LAYOUT_METADATA_UNCERTAIN" in result.codes
    assert not result.document.tables


@pytest.mark.parametrize(
    "raw", ["text[[a,2,3,4]]", "table[[0,0,1001,2]]", "table[[[1,2,3,4]]]", "text[[1,2,3,4]]]"]
)
def test_closed_uncertain_layout(raw):
    ex = extract(raw, model_id="deepseek-ocr:3b")
    assert ex.values == [raw] and ex.layout_uncertain == 1


def test_model_scope_natural_collision_and_no_bbox_grid():
    raw = "table[[1,2,3,4]]\n0012"
    assert extract(raw, model_id="glm-ocr:latest").values == ["table[[1,2,3,4]]", "0012"]
    assert not extract(raw, model_id="deepseek-ocr:3b").document.tables
    natural = "Use table[[1,2,3,4]] in prose"
    assert extract(natural, model_id="deepseek-ocr:3b").values == [natural]
    assert extract(
        "<table><tr><td>text[[1,2,3,4]]</td></tr></table>", model_id="deepseek-ocr:3b"
    ).values == ["text[[1,2,3,4]]"]


def test_islands_keep_values_and_source_table_order():
    md = "| A | B |\n| --- | --- |\n| 001 | 2 |\n\n"
    html = '<table><tr><td rowspan="2">X</td><td>Y</td></tr><tr><td>Z</td></tr></table>'
    bad = '<table><tr><td colspan="100001">Q</td></tr></table>'
    ex = extract(md + html + bad, model_id="glm-ocr:latest")
    assert len(ex.document.tables) == 2
    assert ex.document.tables[0].cells[0].text == "A"
    assert ex.document.tables[1].cells[0].rowspan == 2
    assert ex.lines == ["Q"]
    assert "RESPONSE_STRUCTURE_LIMIT" in ex.codes
    assert "GENERATION_LIMIT_CONFIRMED" not in ex.codes


@pytest.mark.parametrize(
    "raw",
    [
        '<table><tr><td rowspan="2">A</td></tr><tr><td>B',
        '<table><tr><td colspan="0">A</td></tr><tr><td>B',
        "<table><tr><td>A</td><td>B</td></tr><tr><td>C</td></tr></table>",
        "<table><tr><td>A<table><tr><td>B</td></tr></table></td></tr></table>",
    ],
)
def test_no_invented_grid_from_broken_spans_rows_or_nesting(raw):
    result = extract(raw, model_id="glm-ocr:latest")
    assert not result.document.tables
    assert "A" in result.values
    assert "B" in result.values or "C" in result.values


@pytest.mark.parametrize(
    "raw",
    [
        "| A | B |\n| --- | --- | --- |\n| 1 | 2 |",
        "| A | B |\n| --- | --- |\n| 1 |",
        "| A | B |\n| --- | --- |\n| a\\|b | c |",
    ],
)
def test_markdown_delimiters_are_authoritative(raw):
    with pytest.raises(StructureError):
        parse_response(raw)
    assert not extract(raw, model_id="glm-ocr:latest").document.tables


def test_explicit_empty_edge_cells_and_ambiguous_attributes():
    table = parse_response("| | B | |\n| --- | --- | --- |\n| | 0012 | | ").tables[0]
    assert table.columns == 3 and [c.text for c in table.cells] == ["", "B", "", "", "0012", ""]
    with pytest.raises(StructureError):
        parse_response('<table><tr><td colspan="1" colspan="2">A</td></tr></table>')


@pytest.mark.parametrize("family", ["ollama", "chat"])
@pytest.mark.parametrize(
    "reason,state",
    [
        ("length", "GENERATION_LIMIT_CONFIRMED"),
        ("stop", "NORMAL_STOP"),
        (None, "FINISH_UNKNOWN"),
        ("other", "FINISH_UNKNOWN"),
        ([], "FINISH_UNKNOWN"),
    ],
)
def test_api_finish_is_independent_of_parser_suspicions(family, reason, state):
    raw = "<table><tr><td>0012"
    data = (
        {"response": raw, "done": True, "done_reason": reason}
        if family == "ollama"
        else {"choices": [{"message": {"content": raw}, "finish_reason": reason}]}
    )
    metadata = completion_metadata(data, family)
    assert metadata["finish_state"] == state
    result = extract(raw, completion=metadata)
    assert "RESPONSE_INCOMPLETE_SUSPECTED" in result.codes
    assert ("TRUNCATED_CONFIRMED" in result.codes) == (reason == "length")
    assert result.values == ["0012"]
    complete = extract("0012", completion=metadata)
    assert "RESPONSE_INCOMPLETE_SUSPECTED" not in complete.codes


@pytest.mark.parametrize("done", [False, 1, "true", None])
def test_done_false_is_protocol_evidence_not_generation_limit(done):
    metadata = completion_metadata(
        {"response": "0012", "done": done, "done_reason": "stop"}, "ollama"
    )
    assert metadata["transport_incomplete"]
    ex = extract("0012", completion=metadata)
    assert "TRANSPORT_INCOMPLETE" in ex.codes and "TRUNCATED_CONFIRMED" not in ex.codes
    assert "SUSPECTED_TRUNCATION" not in ex.codes


@pytest.mark.parametrize(
    "data,family",
    [
        ({}, "chat"),
        ({"choices": [None]}, "chat"),
        ({"choices": [{"message": None}]}, "chat"),
        ({"response": None}, "ollama"),
        ([], "ollama"),
    ],
)
def test_malformed_protocol(data, family):
    with pytest.raises(OCRServerError, match="MALFORMED_API_RESPONSE"):
        completion_metadata(data, family)


@pytest.mark.parametrize("api", ["direct", "proxy", "chat"])
@pytest.mark.parametrize(
    "status,expected",
    [
        (500, "HTTP_SERVER_5XX"),
        (429, "HTTP_RATE_LIMIT_429"),
        (401, "HTTP_AUTH_401"),
        (403, "HTTP_FORBIDDEN_403"),
        (404, "HTTP_NOT_FOUND_404"),
    ],
)
def test_real_http_budget_no_fanout_and_error_categories(server, api, status, expected):
    url, state = server
    state["status"] = status
    state["body"] = b"{}"
    if api == "direct":
        from localocr.ollama import OllamaClient

        client = OllamaClient(Settings(url=url, retries=2, max_calls=24))
    else:
        cls = OpenWebUIChatProvider if api == "chat" else OpenWebUIOllamaProxyProvider
        client = cls(OpenWebUISettings(url=url, api=api, retries=2, max_calls=24))
    with pytest.raises(OCRServerError) as error:
        recognize_failsoft(Image.new("RGB", (16, 16)), client, None, raw_only=True)
    assert failure_code(error.value, "ocr") == expected
    assert client.http_ocr_calls == (3 if status == 500 else 1)
    assert len(state["requests"]) == client.http_ocr_calls
    assert all(p["model"] == client.settings.model for _, p in state["requests"])


def test_request_cap_max_calls_and_deadline(server):
    url, state = server
    state["status"], state["body"] = 500, b"{}"
    client = OpenWebUIOllamaProxyProvider(OpenWebUISettings(url=url, retries=2, max_calls=1))
    with pytest.raises(OCRServerError, match="IMAGE_REQUEST_BUDGET_EXHAUSTED"):
        recognize_failsoft(Image.new("RGB", (16, 16)), client, None, raw_only=True)
    assert client.http_ocr_calls == 1
    state["delay"] = 0.4
    client = OpenWebUIOllamaProxyProvider(OpenWebUISettings(url=url, retries=2, timeout=1))
    client.image_deadline = 0
    client.request_budget_remaining = 3
    with pytest.raises(OCRServerError, match="TIMEOUT"):
        client.request("POST", "/api/generate", {"images": ["synthetic"]})
    assert getattr(client, "http_ocr_calls", 0) == 0


def test_cancel_real_http_then_next(server):
    url, state = server
    state["delay"] = 1
    cancel = Event()
    client = OpenWebUIOllamaProxyProvider(OpenWebUISettings(url=url, retries=2), cancel)
    timer = Timer(0.15, cancel.set)
    timer.start()
    try:
        with pytest.raises(Cancelled):
            recognize_failsoft(Image.new("RGB", (16, 16)), client, None, raw_only=True)
    finally:
        timer.join()
    assert client.http_ocr_calls <= 1
    state["delay"] = 0
    next_client = OpenWebUIOllamaProxyProvider(client.settings)
    raw, truncated = recognize_failsoft(
        Image.new("RGB", (16, 16)), next_client, None, raw_only=True
    )
    assert raw == state["text"] and not truncated
    assert next_client.http_ocr_calls == 1


def test_service_server_failure_does_not_stop_queue(server, tmp_path, monkeypatch):
    url, state = server
    state["status"], state["body"] = 500, b"{}"

    class Provider(OpenWebUIOllamaProxyProvider):
        def doctor(self):
            self.model_identity = {"digest": "synthetic-fixed"}
            return "ready"

        def request(self, method, endpoint, payload=None):
            try:
                return super().request(method, endpoint, payload)
            finally:
                state["status"], state["body"] = 200, None

    client = Provider(OpenWebUISettings(url=url, retries=2))
    monkeypatch.setattr(service, "create_provider", lambda *a, **kw: client)
    source = tmp_path / "input"
    source.mkdir()
    for name in ("a.png", "b.png"):
        Image.new("RGB", (16, 16)).save(source / name)
    results = service.convert(source, client.settings, tmp_path / "out", fail_soft=True)
    assert [r.status for r in results] == ["FAILED", "REVIEW"]
    assert [r.http_attempts for r in results] == [3, 1]
    assert results[0].request_metrics["http_status"] == 500
    assert not results[0].output and results[1].ocr_contentful and results[1].table_structured


def test_worker_layout_provenance_and_finish_readback(tmp_path):
    case = FIXTURES[0]
    metadata = completion_metadata(
        {"response": case["raw"], "done": True, "done_reason": "length"}, "ollama"
    )
    private = tmp_path / "private"
    result = postprocess.convert_raw(
        case["raw"],
        tmp_path / "result.xlsx",
        model_id=case["model"],
        completion=metadata,
        diagnostics=private,
    )
    outcome = result["outcome"]
    assert outcome["layout_markers_removed"] == 3
    assert outcome["completion"]["finish_state"] == "GENERATION_LIMIT_CONFIRMED"
    assert "TRUNCATED_CONFIRMED" in outcome["error_codes"]
    provenance = json.loads((private / "layout-provenance.json").read_text(encoding="utf-8"))
    assert len(provenance["metadata"]) == 3
    literal = postprocess.convert_raw(
        "В инструкции table[[1,2,3,4]] — пример", tmp_path / "natural.xlsx", model_id=case["model"]
    )
    assert literal["outcome"]["layout_markers_remaining"] == 1
    assert literal["outcome"]["layout_markers_removed"] == 0


def test_parser_limits_are_bounded_and_dont_suppress_neighbor(tmp_path, monkeypatch):
    monkeypatch.setattr(deepseek_layout, "MAX_MARKERS", 2)
    with pytest.raises(ValueError, match="LAYOUT_MARKER_LIMIT"):
        extract("text[[1,2,3,4]]\n" * 3, model_id="deepseek-ocr:3b")
    raw = (
        "<table><tr><td>0012</td></tr></table>"
        + "<table><tr><td>"
        + "<b>" * 129
        + "X"
        + "</b>" * 129
        + "</td></tr></table>"
    )
    ex = extract(raw, model_id="glm-ocr:latest")
    assert ex.document.tables[0].cells[0].text == "0012" and ex.lines == ["X"]
    assert "RESPONSE_STRUCTURE_LIMIT" in ex.codes


def test_legacy_journal_and_parser_prompt_fingerprint(tmp_path):
    source = tmp_path / "input.png"
    Image.new("RGB", (8, 8)).save(source)
    out = tmp_path / "results"
    out.mkdir()
    legacy = {
        "key": "legacy",
        "status": "PARTIAL",
        "result": {"error_codes": ["SUSPECTED_TRUNCATION"]},
    }
    (out / "jobs.jsonl").write_text(json.dumps(legacy) + '\n{"torn":', encoding="utf-8")
    settings = OpenWebUISettings(api="proxy")
    client = OpenWebUIOllamaProxyProvider(settings)
    identity = settings_identity(settings, client)
    with Journal(out) as journal:
        assert journal.records["legacy"] == legacy
        a, _ = journal.key(source, identity)
        b, _ = journal.key(source, {**identity, "pipeline": "evidence-grid-1/stage27-failsoft/1"})
        c, _ = journal.key(source, settings_identity(replace(settings, num_predict=512), client))
        assert len({a, b, c}) == 3
        journal.append({"key": a, "status": "FAILED"})
    with Journal(out) as journal:
        assert journal.records["legacy"] == legacy and a in journal.records
    assert JobResult("old", "PARTIAL").completion == {}


def test_calibration_counts_text_span_and_multiplicity_independently():
    score = runpy.run_path("scripts/replay_stage210.py")["score"]

    gold = {
        "rows": 1,
        "columns": 2,
        "cells": [
            {"row": 0, "column": 0, "rowspan": 1, "colspan": 1, "text": "0012"},
            {"row": 0, "column": 1, "rowspan": 1, "colspan": 1, "text": "A"},
        ],
    }
    exact = score(extract("<table><tr><td>0012</td><td>A</td></tr></table>"), gold)
    assert exact["whole_table_exact"] == 1 and exact["nonempty_text_exact"] == 2
    assert exact["nonempty_row_col_span_text_joint_exact"] == 2 and exact["numeric_exact"] == 1
    unplaced = score(extract("0012\nA\nA"), gold)
    assert unplaced["whole_table_exact"] == 0 and unplaced["nonempty_text_exact"] == 0
    assert unplaced["text_multiset_exact"] == 2 and unplaced["missing_anchors"] == 2
    swapped = score(extract("<table><tr><td>A</td><td>0012</td></tr></table>"), gold)
    assert swapped["nonempty_text_exact"] == 0 and swapped["row_col_span_exact"] == 2


def test_crlf_raw_bytes_survive_http_diagnostics_and_worker(server, tmp_path):
    case = FIXTURES[1]
    url, state = server
    state["text"] = case["raw"]
    client = OpenWebUIOllamaProxyProvider(OpenWebUISettings(url=url, model=case["model"]))
    private = tmp_path / "private"
    raw, truncated = recognize_failsoft(
        Image.new("RGB", (16, 16)),
        client,
        private,
        prompt=failsoft.selected_prompt(client.settings),
        raw_only=True,
    )
    assert raw == case["raw"] and not truncated
    assert (private / "response-1.txt").read_bytes() == raw.encode("utf-8")
    identity = json.loads((private / "response-identity.json").read_text("utf-8"))
    assert (
        identity["sha256"] == hashlib.sha256((private / "response-1.txt").read_bytes()).hexdigest()
    )
    result = postprocess.convert_raw(raw, tmp_path / "crlf.xlsx", model_id=case["model"])
    assert result["tables"] == 1 and result["outcome"]["table_structured"]
    book = load_workbook(tmp_path / "crlf.xlsx")
    assert [c.value for row in book["Таблица 1"] for c in row] == case["values"]
    book.close()


def test_explicit_newline_preserves_already_crlf_response():
    # Reproduce the Windows default TextIO conversion independently of the host OS.
    raw = FIXTURES[1]["raw"]
    translated = io.BytesIO()
    stream = io.TextIOWrapper(translated, encoding="utf-8", newline="\r\n")
    stream.write(raw)
    stream.flush()
    assert b"\r\r\n" in translated.getvalue()
    exact = io.BytesIO()
    stream2 = io.TextIOWrapper(exact, encoding="utf-8", newline="")
    stream2.write(raw)
    stream2.flush()
    assert exact.getvalue() == raw.encode("utf-8")
