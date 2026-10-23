"""Validated tables to atomic OOXML; OCR strings never become formulas."""

import math
import os
import re
import tempfile
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from threading import Event
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import ImageFont

from localocr.domain import TableDocument, ValidationIssue
from localocr.ollama import Cancelled

MIN_COLUMN_WIDTH = 8.0
MAX_COLUMN_WIDTH = 52.0
MAX_ROW_HEIGHT = 150.0
FONT_SIZE = 11


def value(text: str):
    if re.fullmatch(r"-?(?:0|[1-9]\d{0,14})", text):
        return int(text) if text != "-0" else text
    if re.fullmatch(r"-?(?:0|[1-9]\d{0,12})[,.]\d{1,8}", text):
        number = float(text.replace(",", "."))
        if (
            math.isfinite(number)
            and len(text.replace("-", "").replace(",", "").replace(".", "")) <= 15
            and Decimal(str(number)) == Decimal(text.replace(",", "."))
            and not (number == 0 and text.startswith("-"))
        ):
            return number
    return text


@lru_cache(maxsize=1)
def _font_paths() -> tuple[Path, Path] | None:
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [
        (windows / "arial.ttf", windows / "arialbd.ttf"),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ),
    ]
    return next((pair for pair in candidates if all(path.is_file() for path in pair)), None)


@lru_cache(maxsize=64)
def _font(size: int, bold: bool = False):
    pixels = max(1, round(size * 96 / 72))
    paths = _font_paths()
    if paths is not None:
        try:
            return ImageFont.truetype(str(paths[int(bold)]), size=pixels)
        except OSError, ImportError:
            pass
    try:
        return ImageFont.load_default(size=pixels)
    except OSError, TypeError, ImportError:
        return None


def _text_width(text: str, size: int, bold: bool = False) -> float:
    font = _font(size, bold)
    if font is None:
        width = sum(1.0 if ord(char) < 128 else 1.8 for char in text)
        return width * size * 96 / 72 * (0.56 if bold else 0.52)
    try:
        fallback = font.getname()[0] == "Aileron"
        if fallback:
            width = sum(
                float(font.getlength(char)) if ord(char) < 128 else size * 96 / 72 * 0.9
                for char in text
            )
            return width * (1.08 if bold else 1.0)
        return float(font.getlength(text))
    except AttributeError, OSError, ValueError:
        width = sum(1.0 if ord(char) < 128 else 1.8 for char in text)
        return width * size * 96 / 72 * (0.56 if bold else 0.52)


def _display_text(cell) -> str:
    if cell.value is None:
        return ""
    if isinstance(cell.value, float):
        section = cell.number_format.split(";", 1)[0]
        decimals_match = re.search(r"\.([0#]+)", section)
        decimals = len(decimals_match.group(1)) if decimals_match else 0
        number = cell.value * 100 if "%" in section else cell.value
        text = f"{number:,.{decimals}f}" if decimals else f"{number:,.0f}"
        currency = next((symbol for symbol in "$€£₽¥" if symbol in section), "")
        if currency:
            text = currency + text
        if "%" in section:
            text += "%"
        return text
    return str(cell.value)


def _wrapped_lines(text: str, width: float, size: int, bold: bool) -> int:
    paragraphs = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    count = 0
    for paragraph in paragraphs:
        if not paragraph:
            count += 1
            continue
        line = ""
        for word in paragraph.split(" "):
            candidate = word if not line else line + " " + word
            if _text_width(candidate, size, bold) <= width:
                line = candidate
                continue
            if line:
                count += 1
                line = ""
            for char in word:
                candidate = line + char
                if line and _text_width(candidate, size, bold) > width:
                    count += 1
                    line = char
                else:
                    line = candidate
        count += bool(line)
    return max(1, count)


def _rotated_extent(text: str, size: int, bold: bool, rotation: int) -> tuple[float, float]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    line_height = (size * 1.25 + 2.0) * 96 / 72
    if rotation == 255:
        characters = "".join(lines)
        return (
            max((_text_width(char, size, bold) for char in characters), default=0.0),
            max(1, len(characters)) * line_height,
        )
    width = max((_text_width(line, size, bold) for line in lines), default=0.0)
    height = len(lines) * line_height
    angle = math.radians(rotation if rotation <= 90 else 180 - rotation)
    return (
        width * abs(math.cos(angle)) + height * math.sin(angle),
        width * math.sin(angle) + height * abs(math.cos(angle)),
    )


def _autofit_sheet(sheet) -> list[tuple[int, int, str]]:
    columns = sheet.max_column
    rows = sheet.max_row
    widths = [MIN_COLUMN_WIDTH] * columns
    merged = {}
    for area in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = area.bounds
        merged[min_row, min_col] = (max_row - min_row + 1, max_col - min_col + 1)
    digit_width = max(1.0, _text_width("0", FONT_SIZE))
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            row_span, column_span = merged.get((cell.row, cell.column), (1, 1))
            if column_span != 1:
                continue
            text = _display_text(cell)
            size = max(8, min(36, round(cell.font.sz or FONT_SIZE)))
            bold = bool(cell.font.bold)
            rotation = cell.alignment.textRotation
            desired, _ = _rotated_extent(text, size, bold, rotation)
            width = max(MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, desired / digit_width + 2.5))
            widths[cell.column - 1] = max(widths[cell.column - 1], width)
    for (row_number, column_number), (row_span, column_span) in merged.items():
        cell = sheet.cell(row_number, column_number)
        text = _display_text(cell)
        size = max(8, min(36, round(cell.font.sz or FONT_SIZE)))
        bold = bool(cell.font.bold)
        extent, _ = _rotated_extent(text, size, bold, cell.alignment.textRotation)
        desired = extent / digit_width + 2.5
        available = sum(widths[column_number - 1 : column_number + column_span - 1])
        available += max(0, column_span - 1) * 0.7
        extra = min(max(0.0, desired - available), MAX_COLUMN_WIDTH * column_span - available)
        if extra:
            per_column = extra / column_span
            for index in range(column_number - 1, column_number + column_span - 1):
                widths[index] = min(MAX_COLUMN_WIDTH, widths[index] + per_column)
    widths = [math.ceil(width * 10) / 10 for width in widths]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    heights = [15.0] * rows
    overflow = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            row_span, column_span = merged.get((cell.row, cell.column), (1, 1))
            size = max(8, min(36, round(cell.font.sz or FONT_SIZE)))
            bold = bool(cell.font.bold)
            width = sum(widths[cell.column - 1 : cell.column + column_span - 1])
            width += max(0, column_span - 1) * 0.7
            capacity = max(1.0, (width - 2.5) * digit_width)
            text = _display_text(cell)
            rotation = cell.alignment.textRotation
            if rotation:
                _, pixel_height = _rotated_extent(text, size, bold, rotation)
                required = 3.0 + pixel_height * 72 / 96
            else:
                line_count = _wrapped_lines(text, capacity, size, bold)
                required = 3.0 + line_count * (size * 1.25 + 2.0)
            available = sum(heights[cell.row - 1 : cell.row + row_span - 1])
            if required > MAX_ROW_HEIGHT * row_span:
                overflow.append(
                    (cell.row, cell.column, "DISPLAY_LIMIT: текст превышает высоту строки Excel")
                )
            desired = min(required, MAX_ROW_HEIGHT * row_span)
            if desired > available:
                indices = range(cell.row - 1, cell.row + row_span - 1)
                remaining = sum(MAX_ROW_HEIGHT - heights[index] for index in indices)
                extra = desired - available
                for index in indices:
                    heights[index] += extra * (MAX_ROW_HEIGHT - heights[index]) / remaining
    for index, height in enumerate(heights, 1):
        sheet.row_dimensions[index].height = round(height, 1)
    return overflow


def write_workbook(
    document: TableDocument, path: Path, overwrite=False, cancel: Event | None = None
):
    document.validate()

    def original_region(region):
        matrix = document.normalization.get("normalized_to_original")
        if not matrix or not region:
            return ""
        left, top, right, bottom = region
        return str(
            [
                tuple(round(sum(a * b for a, b in zip(row, (x, y, 1))), 3) for row in matrix)
                for x, y in ((left, top), (right, top), (right, bottom), (left, bottom))
            ]
        )

    book = Workbook()
    book.remove(book.active)
    font = _font(FONT_SIZE)
    font_name = font.getname()[0] if font is not None else "Arial"
    if font_name == "Aileron":
        font_name = "Arial"
    book._named_styles["Normal"].font = Font(name=font_name, sz=FONT_SIZE)
    review = []
    border = Border(*([Side(style="thin", color="808080")] * 4))
    for index, table in enumerate(document.tables, 1):
        sheet = book.create_sheet(f"Таблица {index}")
        for cell in table.cells:
            target = sheet.cell(cell.row + 1, cell.column + 1)
            if ILLEGAL_CHARACTERS_RE.search(cell.text) or len(cell.text) > 32767:
                raise ValueError(
                    "Текст ячейки не помещается в Excel или содержит control characters"
                )
            target.value = value(cell.text)
            if isinstance(target.value, int):
                target.number_format = "0"
            if isinstance(target.value, float):
                decimals = len(re.split(r"[,.]", cell.text)[1])
                target.number_format = "0." + "0" * decimals
            if isinstance(target.value, str):
                target.data_type = "s"
            target.font = Font(
                name=font_name,
                sz=FONT_SIZE,
                bold=cell.is_header,
            )
            target.alignment = Alignment(vertical="center", wrap_text=True)
            target.border = border
            if cell.warnings:
                target.fill = PatternFill("solid", fgColor="FFF2CC")
                review.append(
                    (
                        index,
                        cell.row + 1,
                        cell.column + 1,
                        "; ".join(cell.warnings),
                        str(cell.source_region or ""),
                        str(cell.provenance.source_polygon or ""),
                    )
                )
            if cell.rowspan > 1 or cell.colspan > 1:
                sheet.merge_cells(
                    start_row=cell.row + 1,
                    start_column=cell.column + 1,
                    end_row=cell.row + cell.rowspan,
                    end_column=cell.column + cell.colspan,
                )
        for row in sheet:
            for target in row:
                target.border = border
        for row, column, message in _autofit_sheet(sheet):
            issue = ValidationIssue("DISPLAY_LIMIT", message, row - 1, column - 1)
            if issue not in table.issues:
                table.issues.append(issue)
        grid = table.grid()
        headers = 0
        for row in range(table.rows):
            if all(grid[row, column].is_header for column in range(table.columns)):
                headers += 1
            else:
                break
        if 0 < headers < table.rows and all(
            c.row + c.rowspan <= headers for c in table.cells if c.row < headers
        ):
            sheet.freeze_panes = f"A{headers + 1}"
        review.extend(
            (
                index,
                None if issue.row is None else issue.row + 1,
                None if issue.column is None else issue.column + 1,
                issue.message,
                str(issue.source_region or ""),
                original_region(issue.source_region),
            )
            for issue in table.issues
        )
    review.extend(
        (
            None,
            None if issue.row is None else issue.row + 1,
            None if issue.column is None else issue.column + 1,
            issue.message,
            str(issue.source_region or ""),
            original_region(issue.source_region),
        )
        for issue in document.issues
    )
    if review:
        sheet = book.create_sheet("Проверка")
        sheet.append(
            [
                "Таблица",
                "Строка",
                "Столбец",
                "REVIEW",
                "Регион (normalized-image)",
                "Полигон оригинала (x,y)",
            ]
        )
        for row in review:
            sheet.append(list(row))
            for target in sheet[sheet.max_row]:
                if isinstance(target.value, str):
                    target.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(name=font_name, sz=FONT_SIZE, bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name=font_name, sz=FONT_SIZE)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        _autofit_sheet(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".localocr-", suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    try:
        book.save(temporary)
        check = load_workbook(temporary)
        if any(cell.comment is not None for sheet in check for row in sheet for cell in row):
            check.close()
            raise ValueError("В созданной книге неожиданно появились комментарии")
        check.close()
        with ZipFile(temporary) as archive:
            names = archive.namelist()
            if any(
                name.startswith("xl/comments/")
                or name.startswith("xl/drawings/vmlDrawing")
                or (name.endswith(".rels") and b"/comments" in archive.read(name))
                for name in names
            ):
                raise ValueError("OOXML книги содержит комментарии или VML notes")
        if cancel is not None and cancel.is_set():
            raise Cancelled("Обработка отменена")
        if overwrite:
            os.replace(temporary, path)
        else:
            # Exclusive atomic publication requires hardlink support; never copy to final name.
            try:
                os.link(temporary, path)
            except FileExistsError:
                raise
            except OSError as error:
                raise OSError(
                    "Atomic exclusive publication failed: hardlinks unsupported or access denied; "
                    "choose a supported directory (no incomplete output published)"
                ) from error
    finally:
        Path(temporary).unlink(missing_ok=True)
