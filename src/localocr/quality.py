"""Opt-in local qualification. Exported OCR data must never enter a handoff."""

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from localocr.domain import Cell, StructureError, Table


def export(path: Path):
    book = load_workbook(path)
    tables = []
    for sheet in book:
        if sheet.title == "Проверка":
            continue
        covered = set()
        spans = {}
        for merge in sheet.merged_cells.ranges:
            spans[merge.min_row, merge.min_col] = [
                merge.max_row - merge.min_row + 1,
                merge.max_col - merge.min_col + 1,
            ]
            covered.update(
                (r, c)
                for r in range(merge.min_row, merge.max_row + 1)
                for c in range(merge.min_col, merge.max_col + 1)
                if (r, c) != (merge.min_row, merge.min_col)
            )
        cells = []
        for row in sheet:
            for cell in row:
                if (cell.row, cell.column) in covered:
                    continue
                cells.append(
                    dict(
                        row=cell.row,
                        column=cell.column,
                        spans=spans.get((cell.row, cell.column), [1, 1]),
                        value=cell.value,
                        type=cell.data_type,
                        number_format=cell.number_format,
                        text=cell.comment.text if cell.comment else None,
                        review=cell.fill.fgColor.rgb == "00FFF2CC",
                    )
                )
        tables.append(dict(rows=sheet.max_row, columns=sheet.max_column, cells=cells))
    result = dict(
        tables=tables,
        review_count=book["Проверка"].max_row - 1 if "Проверка" in book.sheetnames else 0,
    )
    book.close()
    return result


def template(rows=35, columns=9):
    return dict(
        human_verified=False,
        tables=[
            dict(
                rows=rows,
                columns=columns,
                cells=[
                    dict(row=r, column=c, spans=None, text=None, type=None)
                    for r in range(1, rows + 1)
                    for c in range(1, columns + 1)
                ],
            )
        ],
    )


def compare(actual, gold):
    if not gold.get("human_verified") or any(
        cell.get("text") is None or cell.get("spans") is None or cell.get("type") is None
        for table in gold["tables"]
        for cell in table["cells"]
    ):
        return dict(status="QUALITY_PENDING", reason="Complete independent human gold required")
    try:
        if not gold["tables"]:
            raise StructureError("No gold tables")
        for table in gold["tables"]:
            Table(
                [
                    Cell(c["row"] - 1, c["column"] - 1, c["text"], *c["spans"])
                    for c in table["cells"]
                ],
                table["rows"],
                table["columns"],
            ).grid()
    except StructureError, TypeError, KeyError:
        return dict(status="QUALITY_PENDING", reason="Incomplete/invalid human gold geometry")
    stats = Counter()
    types = {}
    for index in range(max(len(actual["tables"]), len(gold["tables"]))):
        a = actual["tables"][index] if index < len(actual["tables"]) else dict(cells=[])
        g = gold["tables"][index] if index < len(gold["tables"]) else dict(cells=[])
        if (a.get("rows"), a.get("columns")) != (g.get("rows"), g.get("columns")):
            stats["shape_mismatches"] += 1
        ac = {(c["row"], c["column"]): c for c in a["cells"]}
        gc = {(c["row"], c["column"]): c for c in g["cells"]}
        stats["missing"] += len(gc.keys() - ac.keys())
        stats["extra"] += len(ac.keys() - gc.keys())
        for key in gc.keys() & ac.keys():
            expected, observed = gc[key], ac[key]
            bucket = types.setdefault(expected["type"], Counter())
            bucket["checked"] += 1
            if observed.get("type") != expected["type"]:
                stats["type_mismatches"] += 1
            if observed["spans"] != expected["spans"]:
                stats["span_mismatches"] += 1
            text = observed["text"]
            if text is None:
                stats["original_text_unavailable"] += 1
                text = "" if observed["value"] is None else str(observed["value"])
            if text != expected["text"]:
                stats["text_mismatches"] += 1
                bucket["text_mismatches"] += 1
                if not observed["review"]:
                    stats["unflagged_mismatches"] += 1
    return dict(
        status="QUALIFIED_COMPARISON",
        counts=dict(stats),
        by_type=types,
        review_count=actual["review_count"],
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--gold", type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    actual = export(args.xlsx)
    (args.output / "actual.json").write_text(
        json.dumps(actual, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    gold = json.loads(args.gold.read_text(encoding="utf-8")) if args.gold else template()
    if not args.gold:
        (args.output / "manual-gold.json").write_text(json.dumps(gold, indent=2), encoding="utf-8")
    (args.output / "summary.json").write_text(
        json.dumps(compare(actual, gold), indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
