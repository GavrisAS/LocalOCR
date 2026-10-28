"""Opt-in fail-soft adapter. Unverified structure and text require review."""

import os
import re
import tempfile
import unicodedata
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from threading import Event

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from localocr.domain import StructureError, TableDocument
from localocr.excel import write_workbook
from localocr.ollama import Cancelled

VERSION = "stage210-failsoft/2"
MAX_XLSX_ROWS = 1_048_576
MAX_SALVAGE_ROWS = 100_000
MAX_TABLES = 64
MAX_TOTAL_CELLS = 200_000
MAX_RESPONSE = 4_000_000


@dataclass
class Extraction:
    document: TableDocument = field(default_factory=lambda: TableDocument([]))
    lines: list[str] = field(default_factory=list)
    codes: list[str] = field(default_factory=list)
    source_format: str = "plain"
    discarded_markup: int = 0
    readback_nonempty_data_cells: int = 0

    completion: dict = field(default_factory=dict)
    layout_metadata: list[dict] = field(default_factory=list)
    layout_uncertain: int = 0
    layout_metadata_characters: int = 0

    @property
    def values(self):
        return [c.text for t in self.document.tables for c in t.cells if c.text.strip()] + [
            line for line in self.lines if line.strip()
        ]

    def outcome(self):
        structured = any(c.text.strip() for t in self.document.tables for c in t.cells)
        return dict(
            status="FAILED"
            if not self.values
            else "PARTIAL"
            if self.lines or self.codes
            else "REVIEW",
            xlsx_with_data=self.readback_nonempty_data_cells > 0,
            nonempty_data_cells=self.readback_nonempty_data_cells,
            table_structured=structured,
            partial_text_salvage=bool(self.lines),
            no_content=not self.values,
            diagnostic_only=False,
            nonempty_values=len(self.values),
            numeric_values=sum(bool(re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", v)) for v in self.values),
            text_values=sum(
                not bool(re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", v)) for v in self.values
            ),
            error_codes=list(dict.fromkeys(self.codes + ([] if self.values else ["NO_CONTENT"]))),
            source_format=self.source_format,
            discarded_markup=self.discarded_markup,
            coordinates="format-local only; no verified image geometry",
            metrics_version="ocr-content/1",
            ocr_contentful=bool(self.values),
            content_kind=("PARTIAL_STRUCTURED" if self.lines else "STRUCTURED_TABLE")
            if structured
            else "TEXT_ONLY"
            if self.values
            else "NO_CONTENT",
            distinct_content_values=len(set(self.values)),
            content_characters=sum(map(len, self.values)),
            numeric_text_tokens=sum(
                len(re.findall(r"(?<!\w)[+-]?\d+(?:[.,]\d+)?(?!\w)", v)) for v in self.values
            ),
            unplaced_text=len(self.lines),
            layout_markers_removed=len(self.layout_metadata),
            layout_metadata_characters=self.layout_metadata_characters,
            layout_metadata_uncertain=self.layout_uncertain,
            layout_markers_remaining=sum(
                len(re.findall(r"(?:table|text|sub_title)\s*\[\[", v)) for v in self.values
            ),
            completion=self.completion,
        )


class InertText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "iframe", "object"}:
            self.hidden += 1
        elif tag in {"br", "td", "th", "tr", "p", "div", "table"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "iframe", "object"}:
            self.hidden = max(0, self.hidden - 1)
        elif tag in {"td", "th", "tr", "p", "div", "table"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def visible(raw):
    # Reasoning, grounding labels and coordinates are model metadata, not OCR values.
    raw = re.sub(r"<(?:think|analysis|reasoning)\b[^>]*/>", "", raw, flags=re.I)
    raw = re.sub(r"<(think|analysis|reasoning)\b[^>]*>.*?(?:</\1>|\Z)", "", raw, flags=re.I | re.S)
    raw = re.sub(r"<\|ref\|>.*?<\|/ref\|>", "", raw, flags=re.S)
    raw = re.sub(r"<\|det\|>.*?<\|/det\|>", "", raw, flags=re.S)
    raw = re.sub(r"<\|[^\n]*?\|>", "", raw)
    return raw


def control_code(line):
    if not any(
        not char.isspace() and not unicodedata.category(char).startswith("C") for char in line
    ):
        return "CONTROL_ONLY_OUTPUT"
    if re.fullmatch(r"user|assistant|system|developer|analysis|final|tool", line.strip(), re.I):
        return "CONTROL_ONLY_OUTPUT"
    words = line.split()
    if len(words) >= 32 and len(set(words)) == 1 and not words[0].isdecimal():
        return "DEGENERATE_OUTPUT"
    if re.fullmatch(r"(?:(?:None|null|nan)[\s|,;]*){2,}", line, re.I):
        return "DEGENERATE_OUTPUT"
    if (
        len(line) >= 16
        and len(set(line)) == 1
        and (unicodedata.category(line[0]).startswith("P") or line[0] == "№")
    ):
        return "DEGENERATE_OUTPUT"
    return None


def plain_lines(raw, rejected=None):
    cell_error = bool(re.search(r"<(?:td|th)\b", raw, re.I))
    if re.search(r"</?[a-z][^>]*>", raw, re.I):
        parser = InertText()
        parser.feed(raw)
        parser.close()
        raw = "".join(parser.parts)
    result = []
    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped or re.fullmatch(r"```[\w-]*", stripped):
            continue
        if re.fullmatch(r"\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?", line):
            continue
        code = control_code(line)
        if code:
            if rejected is not None:
                rejected.append(code)
            continue
        if cell_error and stripped.casefold() == "error":
            result.append(line)
            continue
        # Conservative recognized control responses, not arbitrary sentences containing 'error'.
        if re.fullmatch(
            r"(?:Table Recognition:|Free OCR\.|(?:<\|grounding\|>)?Convert the document to markdown\.|"
            r"Error|Error: (?:CUDA|HTTP|connection|model|server|request|timeout|failed to load).*|No (?:text|content)(?: (?:found|detected))?\.?|"
            r"I (?:cannot|can't|am unable to) (?:read|extract|process|see).*|"
            r"(?:HTTP )?(?:401 Unauthorized|403 Forbidden|404 Not Found|429 Too Many Requests)|"
            r"(?:HTTP )?5\d{2} (?:Internal Server Error|Bad Gateway|Service Unavailable|Gateway Timeout)|"
            r"(?:CUDA )?out of memory.*|(?:eval|prompt eval|load|total) time\s*=.*|"
            r"\[end of text\]|generation time: [\d.]+\s*(?:s|ms))",
            stripped,
            flags=re.I,
        ):
            continue
        result.append(line)
    return result


def extract(
    raw: str, *, truncated=False, format_hint=None, model_id="", completion=None
) -> Extraction:
    if len(raw) > MAX_RESPONSE:
        raise ValueError("RESPONSE_LIMIT")
    raw = re.sub(r"<(?:script|style|iframe|object|template)\b[^>]*/>", "", raw, flags=re.I)
    raw = re.sub(
        r"<(script|style|iframe|object|template)\b[^>]*>.*?(?:</\1\s*>|\Z)",
        "",
        raw,
        flags=re.I | re.S,
    )
    raw = visible(raw)
    result = Extraction(completion=completion or {})
    if completion and completion.get("finish_state") == "GENERATION_LIMIT_CONFIRMED":
        result.codes.extend(
            ["GENERATION_LIMIT_CONFIRMED", "TRUNCATED_CONFIRMED", "SUSPECTED_TRUNCATION"]
        )
    elif truncated:
        result.codes.extend(["RESPONSE_INCOMPLETE_SUSPECTED", "SUSPECTED_TRUNCATION"])
    if completion and completion.get("transport_incomplete"):
        result.codes.append("TRANSPORT_INCOMPLETE")
    from localocr.deepseek_layout import adapt, is_deepseek

    if is_deepseek(model_id):
        layout = adapt(raw)
        raw = layout.text
        result.layout_metadata = layout.metadata
        result.layout_uncertain = layout.uncertain
        result.layout_metadata_characters = layout.metadata_characters
        if layout.uncertain:
            result.codes.append("LAYOUT_METADATA_UNCERTAIN")
    # A parser suspicion is independent of the API completion state.
    incomplete = (
        any(
            len(re.findall(r"<" + tag + r"\b", raw, re.I))
            != len(re.findall(r"</" + tag + r"\s*>", raw, re.I))
            for tag in ("table", "tr", "td", "th")
        )
        or sum(bool(re.fullmatch(r"\s*```[\w-]*\s*", line)) for line in raw.splitlines()) % 2
    )
    if incomplete:
        result.codes.extend(
            ["RESPONSE_INCOMPLETE_SUSPECTED", "SUSPECTED_INCOMPLETE", "SUSPECTED_TRUNCATION"]
        )
    if not raw.strip():
        return result
    if raw.count("\n") > 100_000:
        raise ValueError("RESPONSE_LINE_LIMIT")
    if format_hint == "tsv":
        # Only an explicit producer contract makes tabs authoritative delimiters.
        rows = [line.split("\t") for line in raw.splitlines()]
        if rows and len(rows[0]) > 1 and len({len(row) for row in rows}) == 1:
            from localocr.domain import Cell, Table

            result.document.tables = [
                Table(
                    [
                        Cell(r, c, value)
                        for r, row in enumerate(rows)
                        for c, value in enumerate(row)
                    ],
                    len(rows),
                    len(rows[0]),
                )
            ]
            result.document.validate()
            result.source_format = "explicit-tsv"
            return result
    from localocr.structured_blocks import extract_islands

    selected = model_id.casefold().split(":", 1)[0].rsplit("/", 1)[-1]
    extract_islands(raw, result, plain_lines, preserve_rows=selected in {"glm-ocr", "deepseek-ocr"})
    if (
        model_id
        and result.lines
        and not result.document.tables
        and "RESPONSE_STRUCTURE_LIMIT" not in result.codes
    ):
        result.codes.append("SOURCE_FORMAT_UNSTRUCTURED")
    # Artifact-bearing tables cannot acquire invented blank anchors when placeholders are removed.
    keep = []
    for table in result.document.tables:

        def table_control(text):
            # Bare role names can be real cell values; deny control envelopes outside tables.
            if text.strip().casefold() in {
                "user",
                "assistant",
                "system",
                "developer",
                "analysis",
                "final",
                "tool",
            }:
                return None
            return control_code(text)

        codes = [table_control(c.text) for c in table.cells if c.text.strip()]
        if any(codes):
            result.codes.extend(code for code in codes if code)
            result.codes.append("UNRESOLVED")
            result.lines.extend(
                c.text for c in table.cells if c.text.strip() and not table_control(c.text)
            )
        else:
            keep.append(table)
    result.document.tables = keep
    # Canonical table bounds/merge validators remain authoritative.
    for table in result.document.tables:
        table.grid()
        for cell in table.cells:
            cell.provenance.coordinate_space = "ocr-format-local"
            cell.provenance.state = "unverified"
    return result


def excel_text(text):
    """Reversible escape for XML-invalid codepoints (and escape-prefix backslashes)."""
    return "".join(
        "\\\\"
        if c == "\\"
        else f"\\u{ord(c):04x}"
        if (ord(c) < 32 and c not in "\t\n\r")
        or 0xD800 <= ord(c) <= 0xDFFF
        or ord(c) in {0xFFFE, 0xFFFF}
        else c
        for c in text
    )


def chunks(text):
    if not text:
        yield ""
        return
    chunk, size = [], 0
    for c in text:
        width = 2 if ord(c) > 0xFFFF else 1
        if size + width > 32767:
            yield "".join(chunk)
            chunk, size = [], 0
        chunk.append(c)
        size += width
    if chunk:
        yield "".join(chunk)


def write_extraction(
    result: Extraction,
    path: Path,
    *,
    cancel: Event | None = None,
    overwrite=False,
    before_publish=None,
):
    """Use canonical writer, add explicit salvage, read back, then publish atomically."""
    if not result.values:
        raise ValueError("NO_CONTENT")
    if cancel is not None and cancel.is_set():
        raise Cancelled("CANCELLED")
    if (
        sum(len(c.text) for t in result.document.tables for c in t.cells)
        + sum(len(line) for line in result.lines)
        > MAX_RESPONSE * 2
    ):
        raise ValueError("EXCEL_TEXT_BUDGET")
    # Excel cannot represent these cells faithfully in one structured anchor.
    keep = []
    for table in result.document.tables:
        if any(len(excel_text(c.text).encode("utf-16-le")) // 2 > 32767 for c in table.cells):
            result.lines.extend(c.text for c in table.cells if c.text.strip())
            result.codes.append("EXCEL_CELL_LIMIT_UNRESOLVED")
        else:
            keep.append(table)
    result.document.tables = keep
    # Bound total workbook allocation before openpyxl creates any cells.
    if len(result.document.tables) > MAX_TABLES:
        raise ValueError("EXCEL_SHEET_LIMIT")
    if sum(t.rows * t.columns for t in result.document.tables) > MAX_TOTAL_CELLS:
        raise ValueError("EXCEL_CELL_BUDGET")
    row_count = 0
    for line in result.lines:
        row_count += sum(1 for _ in chunks(excel_text(line)))
        if row_count > min(MAX_XLSX_ROWS, MAX_SALVAGE_ROWS):
            raise ValueError("EXCEL_ROW_LIMIT")
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".localocr-failsoft-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    temporary = Path(name)
    temporary.unlink()
    try:
        if result.document.tables:
            # Avoid illegal XML chars and preserve lexical values without changing canonical input.
            from copy import deepcopy

            safe = deepcopy(result.document)
            for table in safe.tables:
                for c in table.cells:
                    c.text = excel_text(c.text)
            write_workbook(safe, temporary, cancel=cancel)
            book = load_workbook(temporary)
            for i, table in enumerate(safe.tables, 1):
                for c in table.cells:
                    target = book[f"Таблица {i}"].cell(c.row + 1, c.column + 1)
                    target.value = c.text
                    target.data_type = "s"
        else:
            book = Workbook()
            book.remove(book.active)
        expected = []
        if any(len(list(chunks(excel_text(line)))) > 1 for line in result.lines):
            result.codes.append("EXCEL_CELL_LIMIT_UNRESOLVED")
        if any(excel_text(v) != v for v in result.values):
            result.codes.append("XML_ESCAPED")
        if result.lines:
            sheet = book.create_sheet(
                "Неразмещённые данные" if result.document.tables else "Распознанный текст"
            )
            row_index = 1
            for line in result.lines:
                for chunk in chunks(excel_text(line)):
                    if cancel is not None and cancel.is_set():
                        raise Cancelled("CANCELLED")
                    target = sheet.cell(row_index, 1, chunk)
                    row_index += 1
                    target.data_type = "s"
                    expected.append(chunk)
            sheet.column_dimensions["A"].width = 100
        review = (
            book["Проверка"] if "Проверка" in book.sheetnames else book.create_sheet("Проверка")
        )
        for line in [
            "PARTIAL / REVIEW: значения OCR не проверены человеком.",
            "Координаты таблиц относятся к формату ответа, не доказаны на изображении.",
            "Структура строк salvage неизвестна; одна строка в A, без выдуманных столбцов.",
            "XML-invalid символы: \\uXXXX; исходный обратный слеш удвоен; длинные строки разбиты.",
            "Формат: " + result.source_format,
            "Коды: " + ", ".join(result.codes),
        ]:
            review.append([line])
            review.cell(review.max_row, 1).data_type = "s"
        review.column_dimensions["A"].width = 110
        review.cell(1, 1).font = Font(bold=True, size=16)
        review.cell(1, 1).fill = PatternFill("solid", fgColor="FFF2CC")
        review.row_dimensions[1].height = 32
        book.active = book.sheetnames.index("Проверка")
        book.save(temporary)
        book.close()
        check = load_workbook(temporary)
        try:
            if set(check.sheetnames) != set(book.sheetnames):
                raise ValueError("SHEET_READBACK")
            for i, table in enumerate(result.document.tables, 1):
                expected_merges = {
                    f"{get_column_letter(c.column + 1)}{c.row + 1}:"
                    f"{get_column_letter(c.column + c.colspan)}{c.row + c.rowspan}"
                    for c in table.cells
                    if c.rowspan > 1 or c.colspan > 1
                }
                if {str(r) for r in check[f"Таблица {i}"].merged_cells.ranges} != expected_merges:
                    raise ValueError("MERGE_READBACK")
            readback_count = sum(
                c.value is not None and bool(str(c.value).strip())
                for s in check
                if s.title != "Проверка"
                for row in s
                for c in row
            )
            if not readback_count:
                raise ValueError("NO_CONTENT_READBACK")
            if any(c.data_type == "f" for s in check for row in s for c in row):
                raise ValueError("FORMULA_READBACK")
            if result.lines:
                s = check[
                    "Неразмещённые данные" if result.document.tables else "Распознанный текст"
                ]
                if [(s.cell(i, 1).value or "") for i in range(1, len(expected) + 1)] != expected:
                    raise ValueError("SALVAGE_READBACK")
            for i, table in enumerate(result.document.tables, 1):
                for c in table.cells:
                    if (
                        check[f"Таблица {i}"].cell(c.row + 1, c.column + 1).value or ""
                    ) != excel_text(c.text):
                        raise ValueError("TABLE_READBACK")
        finally:
            check.close()
        if cancel is not None and cancel.is_set():
            raise Cancelled("CANCELLED")
        if before_publish is not None:
            before_publish()
        if cancel is not None and cancel.is_set():
            raise Cancelled("CANCELLED")
        if overwrite:
            os.replace(temporary, path)
        else:
            os.link(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    result.readback_nonempty_data_cells = readback_count
    return result.outcome()


def native_text(raw):
    """Flatten only native OCR content, recursively; never paths/bboxes/metadata."""

    def content(node):
        if isinstance(node, str):
            yield node
        elif isinstance(node, list):
            for child in node:
                yield from content(child)
        elif isinstance(node, dict) and node.get("type") not in {
            "image",
            "image_body",
            "image_path",
        }:
            yield from content(node.get("content"))

    return "\n".join(
        part
        for page in raw.get("pages", [])
        for block in page.get("blocks", [])
        for part in content(block)
    )


@dataclass
class FailSoftDocument(TableDocument):
    unplaced_lines: list[str] = field(default_factory=list)
    error_codes: list[str] = field(default_factory=list)
    source_format: str = "plain"
    readback_nonempty_data_cells: int = 0

    extraction_metrics: dict = field(default_factory=dict)

    def validate(self):
        if self.tables:
            super().validate()
        elif not any(line.strip() for line in self.unplaced_lines):
            raise StructureError("NO_CONTENT")

    def outcome(self):
        original = super().outcome()
        ex = Extraction(
            TableDocument(self.tables, self.issues, self.normalization),
            self.unplaced_lines,
            self.error_codes,
            self.source_format,
        )
        ex.readback_nonempty_data_cells = self.readback_nonempty_data_cells
        state = ex.outcome()
        original.update(self.extraction_metrics)
        original["status"] = state["status"]
        original.update(
            xlsx_with_data=state["xlsx_with_data"],
            table_structured=state["table_structured"],
            partial_text_salvage=state["partial_text_salvage"],
            error_codes=state["error_codes"],
            no_content=state["no_content"],
            nonempty_data_cells=state["nonempty_data_cells"],
            partial_structured=state["table_structured"]
            and bool(self.unplaced_lines or self.error_codes),
        )
        return original


def write_failsoft(document, path, overwrite=False, cancel=None, before_publish=None):
    if not isinstance(document, FailSoftDocument):
        return write_workbook(document, path, overwrite, cancel)
    ex = Extraction(
        TableDocument(document.tables, document.issues, document.normalization),
        document.unplaced_lines,
        document.error_codes,
        document.source_format,
    )
    write_extraction(ex, path, cancel=cancel, overwrite=overwrite, before_publish=before_publish)
    document.tables = ex.document.tables
    document.unplaced_lines = ex.lines
    document.error_codes = ex.codes
    document.readback_nonempty_data_cells = ex.readback_nonempty_data_cells
    state = ex.outcome()
    for name in (
        "ocr_contentful",
        "content_kind",
        "distinct_content_values",
        "content_characters",
        "numeric_text_tokens",
        "unplaced_text",
    ):
        document.extraction_metrics[name] = state[name]


def recognize_failsoft(
    image, client, diagnostics, source_name=None, *, prompt="Table Recognition:", raw_only=False
):
    import base64
    import hashlib
    import time

    from localocr.embedded import EmbeddedOCRProvider
    from localocr.imaging import encode
    from localocr.journal import atomic_json
    from localocr.ollama import OCRServerError
    from localocr.providers import OpenWebUIChatProvider, chat_payload

    started = time.monotonic()
    if isinstance(client, EmbeddedOCRProvider):
        if (
            not client.ready
            or client.process is None
            or client.process.poll() is not None
            or getattr(client, "failsoft_worker_jobs", 0) >= 16
        ):
            client.close()
            client.doctor()
            client.failsoft_worker_jobs = 0
    client.check_cancel()
    before_http = getattr(client, "http_ocr_calls", 0)
    client.image_deadline = started + client.settings.max_image_seconds
    client.request_budget_remaining = min(3, 1 + client.settings.retries, client.settings.max_calls)
    client.last_completion_metadata = {}
    encoded = encode(image)
    if isinstance(client, OpenWebUIChatProvider):
        endpoint = (
            "/v1/chat/completions"
            if isinstance(client, EmbeddedOCRProvider)
            else "/api/chat/completions"
        )
        payload = chat_payload(client.settings, encoded)
        payload["messages"][0]["content"][0]["text"] = prompt
        data = client.request("POST", endpoint, payload)
        try:
            choice = data["choices"][0]
            raw = choice["message"]["content"]

        except (KeyError, IndexError, TypeError) as error:
            raise OCRServerError("MALFORMED_API_RESPONSE") from error
    else:
        endpoint = "/api/generate"
        payload = dict(
            model=client.settings.model,
            prompt=prompt,
            images=[base64.b64encode(encoded).decode("ascii")],
            stream=False,
            options=dict(
                temperature=0,
                seed=0,
                num_ctx=getattr(client.settings, "num_ctx", 32768),
                num_predict=client.settings.num_predict,
            ),
        )
        data = client.request("POST", endpoint, payload)
        raw = data.get("response")

    if not isinstance(raw, str):
        raise OCRServerError("MALFORMED_API_RESPONSE")
    from localocr.response_protocol import completion_metadata

    family = "chat" if isinstance(client, OpenWebUIChatProvider) else "ollama"
    completion = completion_metadata(data, family)
    completion.update(
        prompt_sha256=hashlib.sha256(prompt.encode()).hexdigest(),
        model_id=client.settings.model,
        model_digest=getattr(client, "model_identity", {}).get("digest"),
        generation=payload.get(
            "options", {"max_tokens": client.settings.num_predict, "temperature": 0}
        ),
        timeout=client.settings.timeout,
    )
    client.last_completion_metadata = completion
    truncated = completion["finish_state"] == "GENERATION_LIMIT_CONFIRMED"
    # Raw sidecars remain opt-in. No response text is emitted via network events.
    if diagnostics is not None:
        diagnostics.mkdir(parents=True, exist_ok=True)
        (diagnostics / "response-1.txt").write_text(
            raw, encoding="utf-8", errors="surrogatepass", newline=""
        )
        atomic_json(
            diagnostics / "response-identity.json",
            dict(
                sha256=hashlib.sha256(raw.encode("utf-8", errors="surrogatepass")).hexdigest(),
                endpoint_family=endpoint,
                truncated=truncated,
                completion=completion,
            ),
        )
    client.check_cancel()
    client.failsoft_worker_jobs = getattr(client, "failsoft_worker_jobs", 0) + 1
    client.last_run_metrics = dict(
        calls=1,
        http_attempts=getattr(client, "http_ocr_calls", before_http + 1) - before_http,
        seconds=time.monotonic() - started,
        stage_seconds={},
    )
    if raw_only:
        return raw, truncated
    result = extract(
        raw,
        truncated=truncated,
        completion=completion,
        model_id="" if isinstance(client, EmbeddedOCRProvider) else client.settings.model,
    )
    document = FailSoftDocument(
        result.document.tables,
        result.document.issues,
        image.info.get("localocr_normalization", {}),
        result.lines,
        result.codes,
        result.source_format,
    )
    document.extraction_metrics = result_metrics(result)
    document.validate()
    client.last_run_metrics = dict(
        calls=1,
        http_attempts=getattr(client, "http_ocr_calls", before_http + 1) - before_http,
        seconds=time.monotonic() - started,
        stage_seconds={},
    )
    return document


def write_batch_summary(
    results, output_dir, discovered, cancelled, *, unique_inputs=None, batch=None
):
    import csv
    import math
    from collections import Counter
    from dataclasses import asdict

    from localocr.journal import atomic_json

    all_rows = [asdict(result) for result in results]
    rows = [r for r in all_rows if not r["batch_level"]]
    counts = dict(
        images_seen=discovered,
        terminal_results=len(rows),
        cancelled=cancelled,
        remaining_not_attempted=discovered - len(rows),
        attempted=sum(not r["resumed"] and r["status"] != "SKIPPED" for r in rows),
    )
    for name in ["xlsx_with_data", "table_structured", "partial_text_salvage", "resumed"]:
        counts[name] = sum(bool(r[name]) for r in rows)
    for name in [
        "diagnostic_only",
        "no_content",
        "partial_structured",
        "worker_crashed",
        "timeout",
    ]:
        counts[name] = sum(bool(r[name]) for r in rows)
    for status in ["FAILED", "SKIPPED", "CANCELLED", "BLOCKED_ENV", "BLOCKED_ENDPOINT"]:
        counts[status.lower()] = sum(r["status"] == status for r in rows)
    times = sorted(r["seconds"] for r in rows if not r["resumed"] and r["status"] != "SKIPPED")
    counts.update(
        seen=discovered,
        terminal=len(rows),
        blocked=counts["blocked_env"] + counts["blocked_endpoint"],
        text_salvage=counts["partial_text_salvage"],
        worker_crashes=counts["worker_crashed"],
        p50=times[max(0, math.ceil(len(times) * 0.5) - 1)] if times else None,
        p95=times[max(0, math.ceil(len(times) * 0.95) - 1)] if times else None,
    )
    counts.update(
        {
            "INPUT_TOTAL": discovered,
            "UNIQUE_INPUTS": unique_inputs,
            "ATTEMPTED": counts["attempted"],
            "HTTP_ATTEMPTS": sum(r["http_attempts"] for r in rows),
            "FINISHED": sum(
                r["status"] in {"SUCCESS", "REVIEW", "PARTIAL", "FAILED"} and not r["resumed"]
                for r in rows
            ),
            "CANCELLED": sum(r["status"] == "CANCELLED" for r in rows),
            "SKIPPED": counts["skipped"],
            "RESUMED": counts["resumed"],
            "FAILED": counts["failed"],
            "BLOCKED": (batch or {}).get("blocked", counts["blocked"]),
            "XLSX_WITH_DATA": counts["xlsx_with_data"],
            "TABLE_STRUCTURED": counts["table_structured"],
            "PARTIAL_STRUCTURED": counts["partial_structured"],
            "TEXT_SALVAGE_ONLY": sum(
                r["partial_text_salvage"] and not r["table_structured"] for r in rows
            ),
            "NO_CONTENT": counts["no_content"],
            "DIAGNOSTIC_ONLY": counts["diagnostic_only"],
        }
    )
    counts.update(
        {
            "COMPLETED": sum(r["status"] in {"SUCCESS", "REVIEW", "PARTIAL"} for r in rows),
            "NOT_RUN": max(0, discovered - len(rows)),
            "XLSX_EXISTS": sum(bool(r["output"]) and Path(r["output"]).is_file() for r in rows),
            "XLSX_READBACK_OK": counts["xlsx_with_data"],
            "OCR_CONTENTFUL": sum(r["ocr_contentful"] for r in rows),
            "STRUCTURED_TABLE": counts["table_structured"],
            "TEXT_ONLY": sum(r["content_kind"] == "TEXT_ONLY" for r in rows),
            "DISTINCT_CONTENT_VALUES_PER_DOCUMENT_SUM": sum(
                r["distinct_content_values"] for r in rows
            ),
            "CONTENT_CHARACTERS": sum(r["content_characters"] for r in rows),
            "NUMERIC_TEXT_TOKENS": sum(r["numeric_text_tokens"] for r in rows),
            "LAYOUT_MARKERS_REMOVED": sum(r["layout_markers_removed"] for r in rows),
            "LAYOUT_METADATA_UNCERTAIN": sum(r["layout_metadata_uncertain"] for r in rows),
            "UNPLACED_TEXT": sum(r["unplaced_text"] for r in rows),
            "LAYOUT_MARKERS_REMAINING": sum(r["layout_markers_remaining"] for r in rows),
        }
    )
    atomic_json(
        output_dir / "summary.json",
        dict(
            counts=counts,
            results=all_rows,
            batch=batch,
            error_breakdown=dict(Counter(code for r in rows for code in r["error_codes"])),
            model_backend_breakdown=dict(
                Counter(r["provider"] + "/" + r["model_id"] + "/" + r["backend"] for r in rows)
            ),
            ocr_calls=sum(r["calls"] for r in rows),
            http_attempts=sum(r["http_attempts"] for r in rows),
        ),
    )
    with (output_dir / "summary.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=[
                "source",
                "status",
                "xlsx_with_data",
                "table_structured",
                "partial_text_salvage",
                "resumed",
                "seconds",
                "error_codes",
            ],
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(rows)


def selected_prompt(settings):
    """Strategy keyed only to actual selected ID; never change the selected model."""
    model = getattr(settings, "model", "").casefold()
    return (
        "<|grounding|>Convert the document to markdown."
        if "deepseek-ocr" in model
        else "Table Recognition:"
    )


def failure_code(error, stage):
    from localocr.ollama import OCRServerError

    if str(error) in {
        "NO_CONTENT",
        "WORKER_CRASH",
        "SOURCE_CHANGED",
        "BACKEND_CHANGED",
        "PARSER_TIMEOUT",
        "WRITER_TIMEOUT",
        "RESOURCE_FAILURE",
        "PARSER_FAILURE",
        "WRITER_FAILURE",
        "IO_FAILURE",
        "LAYOUT_MARKER_LIMIT",
        "MALFORMED_API_RESPONSE",
        "TRANSPORT_INCOMPLETE",
    }:
        return str(error)
    if str(error) in {
        "RESPONSE_LIMIT",
        "RESPONSE_LINE_LIMIT",
        "EXCEL_ROW_LIMIT",
        "EXCEL_SHEET_LIMIT",
        "EXCEL_CELL_BUDGET",
        "EXCEL_TEXT_BUDGET",
    }:
        return "RESOURCE_FAILURE"
    if isinstance(error, MemoryError):
        return "RESOURCE_FAILURE"
    if isinstance(error, TimeoutError) or "TIMEOUT" in str(error):
        return "TIMEOUT"
    if isinstance(error, OCRServerError):
        detail = str(error)
        if "TRANSPORT_INCOMPLETE" in detail:
            return "TRANSPORT_INCOMPLETE"
        if "MALFORMED_API_RESPONSE" in detail:
            return "MALFORMED_API_RESPONSE"
        if "INVALID_JSON" in detail or "EMPTY_RESPONSE" in detail:
            return "TRANSPORT_INCOMPLETE"
        if "IMAGE_REQUEST_BUDGET_EXHAUSTED" in detail:
            return "LIMIT_EXCEEDED"
        if stage == "doctor" and ("MODEL_NOT_FOUND" in detail or "не установлена" in detail):
            return "MODEL_UNAVAILABLE"
        if "RESOURCE_FAILURE" in detail or "MEMORY" in detail.upper():
            return "RESOURCE_FAILURE"
        status = re.search(r"HTTP[ _](\d{3})", detail)
        if status:
            return {
                401: "HTTP_AUTH_401",
                403: "HTTP_FORBIDDEN_403",
                404: "HTTP_NOT_FOUND_404",
                429: "HTTP_RATE_LIMIT_429",
            }.get(int(status[1]), "HTTP_SERVER_5XX" if int(status[1]) >= 500 else "NETWORK_FAILURE")
        if "MEMORY" in detail.upper():
            return "RESOURCE_FAILURE"
        if stage == "doctor" and detail.startswith("Embedded:"):
            return "ENGINE_UNAVAILABLE"
        return "NETWORK_FAILURE"
    if stage == "load":
        return "IMAGE_DECODE_FAILURE"
    if stage in {"fingerprint", "xlsx"} or isinstance(error, OSError):
        return "IO_FAILURE"
    if isinstance(error, StructureError):
        return "PARSER_FAILURE"
    return "MODEL_FAILURE"


def result_summary(results):
    """Shared GUI/CLI wording for orthogonal, readback-qualified data metrics."""
    return (
        f"XLSX с данными: {sum(r.xlsx_with_data for r in results)} · "
        f"Структура ответа: {sum(r.table_structured for r in results)} · "
        f"Сохранённый текст: {sum(r.partial_text_salvage for r in results)} · "
        f"Непустых ячеек: {sum(r.nonempty_data_cells for r in results)}"
    )


def result_metrics(result):
    state = result.outcome()
    return {
        k: state[k]
        for k in (
            "metrics_version",
            "ocr_contentful",
            "content_kind",
            "distinct_content_values",
            "content_characters",
            "numeric_text_tokens",
            "unplaced_text",
            "layout_markers_removed",
            "layout_metadata_characters",
            "layout_metadata_uncertain",
            "layout_markers_remaining",
            "completion",
        )
    }
