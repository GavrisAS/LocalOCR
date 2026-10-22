"""Explicit dev/hardware qualification; synthetic inputs only, no downloads.

Run with uv run python scripts/qualify_gpu.py --engine ... --model ... --mmproj ...
--output .evidence/gpu/report.json. Runs the same F16 pair in CPU/GPU/Auto.
"""

import argparse
import hashlib
import json
import sys
import time
from pathlib import Path

from openpyxl import load_workbook
from qualify_embedded import observation_texts, synthetic

from localocr.embedded import EmbeddedOCRProvider, EmbeddedSettings
from localocr.excel import write_workbook
from localocr.imaging import load_image
from localocr.service import recognize_document


def digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def qualify(engine: Path, model: Path, mmproj: Path, output: Path):
    output.parent.mkdir(parents=True, exist_ok=True)
    lock = json.loads(Path(__file__).with_name("embedded-lock.json").read_text())
    pair = lock["models"]["f16"]
    report = {
        "status": "FAIL",
        "platform": sys.platform,
        "lock_sha256": digest(Path(__file__).with_name("embedded-lock.json")),
        "engine_manifest_sha256": None,
        "manifest_note": "Custom engine qualification: binary hash recorded; installed Portable manifest is qualified separately on Windows",
        "engine_sha256": digest(engine),
        "model_sha256": digest(model),
        "mmproj_sha256": digest(mmproj),
        "runs": [],
    }
    if (
        report["model_sha256"] != pair["model"]["sha256"]
        or report["mmproj_sha256"] != pair["mmproj"]["sha256"]
    ):
        raise ValueError("Qualification requires the pinned F16 pair")
    try:
        for mode in ("cpu", "gpu", "auto"):
            settings = EmbeddedSettings(
                engine=engine, model_path=model, mmproj_path=mmproj, device_mode=mode, timeout=600
            )
            run = {"requested_mode": mode, "cases": []}
            report["runs"].append(run)
            with EmbeddedOCRProvider(settings) as client:
                run["runtime"] = client.doctor()
                run["actual_backend"] = client.active_candidate.backend
                run["device"] = client.device_name
                run["offload_evidence"] = client.offload
                run["load_seconds"] = client.load_seconds
                for rows in (3, 30):
                    image = output.parent / f"synthetic-{mode}-{rows}.png"
                    gold = synthetic(image, rows)
                    started = time.monotonic()
                    diagnostics = output.parent / "diagnostics" / image.stem
                    document = recognize_document(load_image(image), client, diagnostics)
                    seconds = time.monotonic() - started
                    target = image.with_suffix(".xlsx")
                    write_workbook(document, target, overwrite=True)
                    book = load_workbook(target)
                    sheet = book["Таблица 1"]
                    observed = observation_texts(diagnostics / "observations.json", rows, 3)
                    if any(
                        cell.comment is not None
                        for worksheet in book.worksheets
                        for row in worksheet.iter_rows()
                        for cell in row
                    ):
                        raise ValueError("New GPU qualification XLSX contains comments")
                    types = [
                        [sheet.cell(r + 1, c + 1).data_type for c in range(3)] for r in range(rows)
                    ]
                    values = [
                        [sheet.cell(r + 1, c + 1).value for c in range(3)] for r in range(rows)
                    ]
                    values_match = values == [[row[0], int(row[1]), float(row[2])] for row in gold]
                    shape = [sheet.max_row, sheet.max_column]
                    book.close()
                    case = {
                        "rows": rows,
                        "shape": shape,
                        "ocr_pipeline_seconds": seconds,
                        "expected": gold,
                        "observed": observed,
                        "cell_types": types,
                        "values_match": values_match,
                        "inference_seconds": list(client.inference_seconds),
                        "text_mismatches": sum(
                            a != b for er, ar in zip(gold, observed) for a, b in zip(er, ar)
                        ),
                    }
                    run["cases"].append(case)
                    if (
                        not values_match
                        or case["text_mismatches"]
                        or shape != [rows, 3]
                        or any(t != ["s", "n", "n"] for t in types)
                    ):
                        raise ValueError(f"Synthetic gold mismatch: {mode} {rows}")
                if mode in ("gpu", "auto") and not client.offload:
                    raise ValueError("Hardware qualification requires confirmed GPU")
        report["status"] = "GPU_HARDWARE_PASS"
    except Exception as error:
        report["error"] = f"{type(error).__name__}: {error}"
        raise
    finally:
        output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("engine", "model", "mmproj", "output"):
        parser.add_argument("--" + name, required=True, type=Path)
    args = parser.parse_args()
    qualify(
        args.engine.resolve(), args.model.resolve(), args.mmproj.resolve(), args.output.resolve()
    )
