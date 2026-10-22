"""Real synthetic vision → canonical → XLSX qualification, never a health-only PASS."""

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from contextlib import nullcontext
from pathlib import Path

from frozen_smoke import isolated_environment
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


def synthetic(path: Path, rows: int):
    scale_x, scale_y = 220, 76
    image = Image.new("RGB", (3 * scale_x + 41, rows * scale_y + 41), "white")
    draw = ImageDraw.Draw(image)
    font_path = (
        "C:/Windows/Fonts/arial.ttf"
        if os.name == "nt"
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    )
    font = ImageFont.truetype(font_path, 26)
    expected = []
    for row in range(rows):
        values = [f"R{row + 1:02}", str(row + 10), f"{row + 1}.25"]
        expected.append(values)
        for column, text in enumerate(values):
            draw.text((35 + column * scale_x, 42 + row * scale_y), text, font=font, fill="black")
    for column in range(4):
        draw.line(
            (20 + column * scale_x, 20, 20 + column * scale_x, 20 + rows * scale_y),
            fill="black",
            width=2,
        )
    for row in range(rows + 1):
        draw.line(
            (20, 20 + row * scale_y, 20 + 3 * scale_x, 20 + row * scale_y), fill="black", width=2
        )
    image.save(path)
    return expected


def frozen_batch_metadata(stdout: str, names: list[str]):
    startup = [float(value) for value in re.findall(r"model load=([0-9.]+)s", stdout)]
    backends = re.findall(r"Embedded F16 · (CPU|CUDA|VULKAN) ·", stdout)
    if len(startup) != 1 or backends != ["CPU"]:
        raise ValueError("Frozen CPU batch must load exactly one owned CPU engine")
    cases = {}
    for name in names:
        marker = f"] {name}\n"
        begin = stdout.find(marker)
        end = stdout.find(f"] {name}: ", begin)
        if begin < 0 or end < 0:
            raise ValueError(f"Missing frozen batch progress: {name}")
        trace = stdout[begin + len(marker) : end]
        requests = [
            float(value) for value in re.findall(r"Embedded recognize latency=([0-9.]+)s", trace)
        ]
        if not requests:
            raise ValueError(f"Missing actual OCR latency: {name}")
        cases[name] = {"ocr_request_seconds": requests, "technical_trace": trace[-20000:]}
    return {
        "startup_seconds": startup[0],
        "actual_backend": "CPU",
        "shared_engine": True,
        "cases": cases,
    }


def observation_texts(path: Path, rows: int, columns: int) -> list[list[str | None]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    tables = data.get("tables")
    if data.get("schema_version") != 1 or not isinstance(tables, list) or len(tables) != 1:
        raise ValueError("Expected one schema-v1 OCR observations table")
    table = tables[0]
    if (table.get("rows"), table.get("columns")) != (rows, columns):
        raise ValueError("OCR observations dimensions differ from synthetic source")
    result = [[None for _ in range(columns)] for _ in range(rows)]
    for cell in table.get("cells", []):
        row, column = cell.get("row"), cell.get("column")
        if (
            not isinstance(row, int)
            or not isinstance(column, int)
            or not 0 <= row < rows
            or not 0 <= column < columns
            or result[row][column] is not None
            or not isinstance(cell.get("text"), str)
        ):
            raise ValueError("Invalid or duplicate OCR observation coordinate")
        result[row][column] = cell["text"]
    return result


def qualify_once(root: Path, output: Path, sha: str, isolated: bool, relocated=False):
    evidence = {
        "status": "FAIL",
        "platform": sys.platform,
        "precision": "f16",
        "build_commit": sha,
        "network_isolated": isolated,
        "cases": [],
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    manifest = root / "embedded/manifest.json"
    evidence["manifest_sha256"] = hashlib.sha256(manifest.read_bytes()).hexdigest()
    evidence["upstream"] = json.loads(manifest.read_text())["upstream"]
    try:
        # Full relocation, including DLLs and models; original installation is not used.
        context = (
            nullcontext(root.parent)
            if relocated
            else tempfile.TemporaryDirectory(
                prefix="localocr-relocated-", ignore_cleanup_errors=True
            )
        )
        with context as directory:
            moved = root if relocated else Path(directory) / "LocalOCR"
            if not relocated:
                shutil.copytree(root, moved)
            exe = moved / "LocalOCR-CLI/LocalOCR-CLI.exe"
            evidence["device_checks"] = []
            for mode in ("cpu", "auto", "gpu"):
                check = subprocess.run(
                    [str(exe), "doctor", "--provider", "embedded", "--device", mode],
                    text=True,
                    encoding="utf-8",
                    capture_output=True,
                    timeout=600,
                    env=isolated_environment(),
                )
                evidence["device_checks"].append(
                    {
                        "mode": mode,
                        "exit_code": check.returncode,
                        "stdout": check.stdout,
                        "stderr": check.stderr,
                    }
                )
                if check.returncode != (3 if mode == "gpu" else 0):
                    raise ValueError(
                        f"CPU/no-GPU runner device policy failed: {mode}: {check.stdout} {check.stderr}"
                    )
            with tempfile.TemporaryDirectory(prefix="synthetic-batch-", dir=directory) as inputs:
                input_dir = Path(inputs)
                gold = {
                    rows: synthetic(input_dir / f"synthetic-{rows}.png", rows) for rows in (3, 30)
                }
                started = time.monotonic()
                result = subprocess.run(
                    [
                        str(exe),
                        "convert",
                        str(input_dir),
                        "--provider",
                        "embedded",
                        "--precision",
                        "f16",
                        "--timeout",
                        "600",
                        "--device",
                        "cpu",
                        "--debug-http",
                        "--diagnostics",
                    ],
                    text=True,
                    encoding="utf-8",
                    capture_output=True,
                    timeout=1800,
                    env=isolated_environment(),
                )
                evidence["batch_seconds"] = time.monotonic() - started
                if result.returncode not in {0, 1}:
                    raise ValueError(f"Frozen vision batch failed: {result.stdout} {result.stderr}")
                metadata = frozen_batch_metadata(
                    result.stdout, [f"synthetic-{rows}.png" for rows in gold]
                )
                evidence["batch_startup_seconds"] = metadata["startup_seconds"]
                evidence["shared_engine"] = metadata["shared_engine"]
                reports = list((input_dir / "results").glob("report-*.json"))
                if len(reports) != 1:
                    raise ValueError("Expected one real frozen batch report")
                jobs = {
                    item["source"]: item
                    for item in json.loads(reports[0].read_text(encoding="utf-8"))["results"]
                }
                for rows, expected in gold.items():
                    image = input_dir / f"synthetic-{rows}.png"
                    xlsx = input_dir / "results" / (image.stem + ".xlsx")
                    if not xlsx.is_file() or jobs.get(image.name, {}).get("status") not in {
                        "SUCCESS",
                        "REVIEW",
                    }:
                        raise ValueError(
                            f"Frozen vision pipeline failed ({rows} rows): {result.stdout} {result.stderr}"
                        )
                    book = load_workbook(xlsx)
                    sheet = book["Таблица 1"]
                    if any(
                        cell.comment is not None
                        for worksheet in book.worksheets
                        for row in worksheet.iter_rows()
                        for cell in row
                    ):
                        raise ValueError("New frozen XLSX unexpectedly contains comments")
                    shape = [sheet.max_row, sheet.max_column]
                    values = [
                        [sheet.cell(r + 1, c + 1).value for c in range(3)] for r in range(rows)
                    ]
                    types = [
                        [sheet.cell(r + 1, c + 1).data_type for c in range(3)] for r in range(rows)
                    ]
                    expected_values = [[row[0], int(row[1]), float(row[2])] for row in expected]
                    book.close()
                    observed = observation_texts(
                        moved / "diagnostics" / image.stem / "observations.json", rows, 3
                    )
                    mismatches = sum(
                        a != b for er, ar in zip(expected, observed) for a, b in zip(er, ar)
                    )
                    evidence["cases"].append(
                        {
                            "rows": rows,
                            "requested_mode": "cpu",
                            "actual_backend": "CPU",
                            "pipeline_seconds": jobs[image.name]["seconds"],
                            **metadata["cases"][image.name],
                            "shape": shape,
                            "text_mismatches": mismatches,
                            "values_match": values == expected_values,
                            "cell_types": types,
                            "xlsx_sha256": hashlib.sha256(xlsx.read_bytes()).hexdigest(),
                            "exit_code": result.returncode,
                        }
                    )
                    if (
                        shape != [rows, 3]
                        or mismatches
                        or values != expected_values
                        or any(t != ["s", "n", "n"] for t in types)
                    ):
                        raise ValueError("Synthetic OCR cells/shape differ from known source")
            evidence["status"] = "PASS" if sys.platform == "win32" and isolated else "PARTIAL"
    except Exception as error:
        evidence["error"] = str(error)
        raise
    finally:
        output.write_text(json.dumps(evidence, indent=2, ensure_ascii=False), encoding="utf-8")
    if evidence["status"] != "PASS":
        raise ValueError("Windows/network isolation qualification remains PENDING")


def qualify(root: Path, output: Path, sha: str, isolated: bool, relocated=False):
    prior: list[dict[str, object]] = []
    for attempt in (1, 2):
        try:
            qualify_once(root, output, sha, isolated, relocated)
        except Exception as error:
            report: dict[str, object] = (
                json.loads(output.read_text(encoding="utf-8"))
                if output.is_file()
                else {"status": "FAIL", "error": str(error)}
            )
            report["qualification_attempts"] = attempt
            report["prior_startup_failures"] = list(prior)
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
            if attempt == 1 and "ENGINE_STARTUP_TIMEOUT" in str(error):
                prior.append(report)
                continue
            raise
        else:
            report = json.loads(output.read_text(encoding="utf-8"))
            report["qualification_attempts"] = attempt
            report["prior_startup_failures"] = prior
            output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
            return


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--sha", required=True)
    parser.add_argument("--network-isolated", action="store_true")
    parser.add_argument("--relocated", action="store_true")
    args = parser.parse_args()
    qualify(args.root.resolve(), args.output, args.sha, args.network_isolated, args.relocated)
