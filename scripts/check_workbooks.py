"""Verify actual OOXML values/types/merges against private canonical observations."""

import argparse
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    records = json.loads((args.run / "private-results.json").read_text())
    summaries = []
    for record in records:
        target = args.run / f"{record['case']}.xlsx"
        if record["status"] == "FAILED":
            summaries.append(
                {"case": record["case"], "failed_without_workbook": not target.exists()}
            )
            continue
        observed = json.loads((args.run / record["case"] / "observations.json").read_text())
        book = load_workbook(target)
        checked = 0
        good = True
        for index, table in enumerate(observed["tables"], 1):
            sheet = book[f"Таблица {index}"]
            merges = set()
            for cell in table["cells"]:
                actual = sheet.cell(cell["row"] + 1, cell["column"] + 1)
                text = cell["text"]
                if not text:
                    good &= actual.value is None
                elif isinstance(actual.value, str):
                    good &= actual.value == text and actual.data_type == "s"
                elif isinstance(actual.value, (int, float)):
                    try:
                        good &= Decimal(str(actual.value)) == Decimal(text.replace(",", "."))
                    except ArithmeticError:
                        good = False
                    good &= actual.data_type == "n"
                    if isinstance(actual.value, float):
                        good &= actual.number_format == "0." + "0" * len(
                            text.replace(",", ".").split(".")[1]
                        )
                else:
                    good = False
                if cell.get("provenance", {}).get("state") == "unresolved":
                    good &= actual.value is None
                if cell["rowspan"] > 1 or cell["colspan"] > 1:
                    begin = f"{get_column_letter(cell['column'] + 1)}{cell['row'] + 1}"
                    end = f"{get_column_letter(cell['column'] + cell['colspan'])}{cell['row'] + cell['rowspan']}"
                    merges.add(f"{begin}:{end}")
                checked += 1
            good &= {str(merge) for merge in sheet.merged_cells.ranges} == merges
        good &= all(
            cell.comment is None and cell.data_type != "f"
            for sheet in book
            for row in sheet
            for cell in row
        )
        book.close()
        summaries.append(
            {
                "case": record["case"],
                "actual_cells_checked": checked,
                "values_types_formats_merges_no_comments_no_formulas": good,
            }
        )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps({"mode": "OOXML_FIDELITY_NOT_OCR_ACCURACY", "cases": summaries}, indent=2)
    )
    if not all(
        r.get(
            "failed_without_workbook",
            r.get("values_types_formats_merges_no_comments_no_formulas", False),
        )
        for r in summaries
    ):
        raise SystemExit("Workbook fidelity failure; inspect private outputs")
    print(f"Actual workbook fidelity: {len(summaries)} outcomes checked")


if __name__ == "__main__":
    main()
