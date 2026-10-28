"""Bounded synthetic evaluation gates; only aggregate evidence leaves the runner."""

import argparse
import base64
import hashlib
import io
import json
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Event
from typing import Any

from frozen_smoke import isolated_environment
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from verify_release import inspect_zip

from localocr import postprocess, runtime
from localocr.ollama import Cancelled


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def qualify(root, version, commit):
    checks = {}
    env = isolated_environment()
    env["QT_QPA_PLATFORM"] = "offscreen"
    with tempfile.TemporaryDirectory(prefix="LocalOCR Русский путь ") as name:
        work = Path(name)
        moved = work / "Распаковано с пробелами" / "LocalOCR"
        shutil.copytree(root, moved)
        cli = moved / "LocalOCR-CLI/LocalOCR-CLI.exe"
        gui = moved / "LocalOCR.exe"
        cwd = work / "Произвольный CWD"
        cwd.mkdir()
        for key in ("APPDATA", "LOCALAPPDATA", "XDG_CONFIG_HOME"):
            env[key] = str(work / key)
        for full in (False, True):
            subprocess.run(
                [
                    str(gui),
                    *(["--full-gui"] if full else []),
                    "--smoke-test",
                    "--expected-version",
                    version,
                    "--expected-commit",
                    commit,
                ],
                cwd=cwd,
                env=env,
                check=True,
                timeout=30,
            )
        assert json.loads((moved / "config.json").read_text("utf-8"))["fail_soft"] is False
        checks["relocated_simple_full_gui"] = "PASS_STARTUP_SETTINGS_FAILSOFT_CONTROLS"
        for flag, expected in (("--version", version), ("--build-commit", commit)):
            assert (
                subprocess.check_output([str(cli), flag], cwd=cwd, env=env, timeout=30)
                .decode()
                .strip()
                == expected
            )
        checks["frozen_identity_without_python_path"] = "PASS"
        fixture = work / "Синтетический вход.png"
        image = Image.new("RGB", (320, 120), "white")
        ImageDraw.Draw(image).text((10, 10), "Synthetic OCR 001 42", fill="black")
        image.save(fixture)
        state: dict[str, Any] = {"raw": "", "posts": 0, "expected_token": None, "errors": []}

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, format, *args):
                pass

            def respond(self, body):
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(body).encode())

            def do_GET(self):
                self.respond(
                    {
                        "models": [
                            {
                                "name": state.get("model", "glm-ocr:latest"),
                                "digest": "synthetic-fixed",
                            }
                        ],
                        "data": [{"id": state.get("model", "glm-ocr:latest")}],
                    }
                )

            def do_POST(self):
                try:
                    assert self.headers.get("Authorization") == state["expected_token"]
                    payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                    assert payload["stream"] is False
                    expected_prompt = (
                        "<|grounding|>Convert the document to markdown."
                        if payload["model"].startswith("deepseek-ocr:")
                        else "Table Recognition:"
                    )
                    if self.path == "/api/chat/completions":
                        data = payload["messages"][0]["content"]
                        assert (
                            next(p["text"] for p in data if p["type"] == "text") == expected_prompt
                        )
                        encoded = next(
                            p["image_url"]["url"].split(",", 1)[1]
                            for p in data
                            if p["type"] == "image_url"
                        )
                    else:
                        assert payload["prompt"] == expected_prompt
                        encoded = payload["images"][0]
                    decoded = Image.open(io.BytesIO(base64.b64decode(encoded, validate=True)))
                    decoded.verify()
                    state["posts"] += 1
                    if state.get("delay"):
                        time.sleep(state["delay"])
                    if state.get("malformed"):
                        self.respond({"choices": [{}], "response": None})
                        return
                    body = (
                        {
                            "choices": [
                                {
                                    "message": {"content": state["raw"]},
                                    "finish_reason": state.get("reason", "stop"),
                                }
                            ]
                        }
                        if self.path == "/api/chat/completions"
                        else {
                            "response": state["raw"],
                            "done": True,
                            "done_reason": state.get("reason", "stop"),
                        }
                    )
                    self.respond(body)
                except BrokenPipeError, ConnectionResetError:
                    pass
                except Exception as error:
                    state["errors"].append(type(error).__name__)
                    self.send_error(500)

        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            cases = [
                ("plain", "001\n=1+1\nx\x00y\n" + "x" * 33000, "ollama", None),
                ("html", "<table><tr><td>001</td></tr><tr><td>truncated", "proxy", None),
                ("markdown", "| A | B |\n| --- | --- |\n| 001 |", "chat", None),
                ("token", "synthetic value", "chat", "synthetic-evaluation-only"),
                ("empty", "<think>no observations</think>", "ollama", None),
            ]
            fixtures = json.loads(Path("tests/fixtures/stage210.json").read_text("utf-8"))
            for fixture_case in (fixtures[0], fixtures[1], fixtures[6], fixtures[7]):
                for api in ("proxy", "chat"):
                    cases.append((fixture_case["id"] + "-" + api, fixture_case["raw"], api, None))
            cases.extend(
                [
                    ("confirmed-length", "0012", "proxy", None),
                    ("malformed-api", "unused", "chat", None),
                    ("timeout", "unused", "proxy", None),
                ]
            )
            for label, raw, api, token in cases:
                state["raw"] = raw
                fixture_case = next((c for c in fixtures if label.startswith(c["id"] + "-")), None)
                state["model"] = fixture_case["model"] if fixture_case else "glm-ocr:latest"
                state["reason"] = "length" if label == "confirmed-length" else "stop"
                state["malformed"] = label == "malformed-api"
                state["delay"] = 2 if label == "timeout" else 0
                before_posts = state["posts"]
                state["expected_token"] = "Bearer " + token if token else None
                out = work / label
                command = [
                    str(cli),
                    "convert",
                    str(fixture),
                    "--model",
                    state["model"],
                    "--timeout",
                    "1" if label == "timeout" else "30",
                    "--fail-soft",
                    "--output-dir",
                    str(out),
                ]
                url = f"http://127.0.0.1:{server.server_port}"
                command += (
                    ["--ollama-url", url]
                    if api == "ollama"
                    else ["--provider", "openwebui", "--url", url, "--openwebui-api", api]
                )
                if token:
                    command += ["--api-token", token]
                p = subprocess.run(command, cwd=cwd, env=env, capture_output=True, timeout=60)
                assert p.returncode == 1, (label, p.returncode)
                assert not state["errors"], state["errors"]
                assert 1 <= state["posts"] - before_posts <= 3
                report = json.loads(max(out.glob("report-*.json")).read_text("utf-8"))
                item = report["results"][0] if isinstance(report, dict) else report[0]
                target = out / (fixture.stem + ".xlsx")
                if label in {"empty", "malformed-api", "timeout"}:
                    if label == "malformed-api":
                        assert "MALFORMED_API_RESPONSE" in item["error_codes"]
                    if label == "timeout":
                        assert "TIMEOUT" in item["error_codes"]
                    assert item["status"] == "FAILED" and not item["xlsx_with_data"]
                    assert not target.exists()
                    checks[label + "_no_false_yield"] = "PASS"
                    if label == "timeout":
                        time.sleep(1.2)
                    continue
                assert item["status"] in {"PARTIAL", "REVIEW"} and item["xlsx_with_data"]
                book = load_workbook(target)
                assert "Проверка" in book.sheetnames
                assert any(c.value for s in book if s.title != "Проверка" for row in s for c in row)
                assert all(c.data_type != "f" for s in book for row in s for c in row)
                if fixture_case:
                    values = [
                        c.value
                        for s in book
                        if s.title != "Проверка"
                        for row in s
                        for c in row
                        if c.value
                    ]
                    assert sorted(values) == sorted(fixture_case["values"]), (
                        label,
                        len(values),
                        len(fixture_case["values"]),
                    )
                    assert item["tables"] == fixture_case["tables"]
                    assert item["layout_markers_removed"] == fixture_case["removed"]
                if label == "confirmed-length":
                    assert item["completion"]["finish_state"] == "GENERATION_LIMIT_CONFIRMED"
                    assert "TRUNCATED_CONFIRMED" in item["error_codes"]
                book.close()
                checks[label + "_multimodal_readback"] = "PASS"
                if label == "plain":
                    before, posts = digest(target), state["posts"]
                    again = subprocess.run(
                        command, cwd=cwd, env=env, capture_output=True, timeout=60
                    )
                    assert (
                        again.returncode == 1
                        and state["posts"] == posts
                        and digest(target) == before
                    )
                    report = json.loads(max(out.glob("report-*.json")).read_text("utf-8"))
                    assert (
                        report[0]["resumed"]
                        if isinstance(report, list)
                        else report["results"][0]["resumed"]
                    )
                    # A foreign existing output with no journal is never overwritten.
                    foreign = work / "foreign"
                    foreign.mkdir()
                    existing = foreign / target.name
                    existing.write_bytes(b"synthetic existing user file")
                    command[command.index(str(out))] = str(foreign)
                    subprocess.run(
                        command, cwd=cwd, env=env, capture_output=True, timeout=60, check=True
                    )
                    assert (
                        existing.read_bytes() == b"synthetic existing user file"
                        and state["posts"] == posts
                    )
                    checks["resume_no_overwrite"] = "PASS"
        finally:
            server.shutdown()
            server.server_close()
        # Actual frozen postprocessor, source parent: exercise parent-only late cancellation.
        old_root, old_argv = runtime.ROOT_OVERRIDE, postprocess.worker_argv
        runtime.ROOT_OVERRIDE = moved
        setattr(
            postprocess,
            "worker_argv",
            lambda s, t: [str(cli), "--localocr-postprocess", str(s), str(t)],
        )
        try:
            target = work / "atomic.xlsx"
            target.write_bytes(b"previous synthetic file")
            cancel = Event()
            try:
                postprocess.convert_raw(
                    "001\n=1+1", target, overwrite=True, cancel=cancel, before_publish=cancel.set
                )
            except Cancelled:
                pass
            else:
                raise AssertionError("Cancellation published a workbook")
            assert target.read_bytes() == b"previous synthetic file"
            postprocess.convert_raw("next observed value", target, overwrite=True)
            book = load_workbook(target)
            assert book["Распознанный текст"]["A1"].value == "next observed value"
            book.close()
            checks["frozen_worker_late_cancel_then_next_readback"] = "PASS"
            old_parser, old_writer, old_launch = (
                postprocess.PARSER_SECONDS,
                postprocess.WRITER_SECONDS,
                postprocess.launch_owned,
            )
            try:
                for fault in ("parser", "writer", "crash", "cancel"):
                    fault_target = work / (fault + ".xlsx")
                    postprocess.PARSER_SECONDS = 0.001 if fault == "parser" else 30.0
                    postprocess.WRITER_SECONDS = 0.001 if fault == "writer" else 30.0
                    fault_cancel = Event()
                    timer = threading.Timer(0.05, fault_cancel.set) if fault == "cancel" else None

                    def launch(*args, **kwargs):
                        child = old_launch(*args, **kwargs)
                        if fault == "crash":
                            # Job close is normal cleanup and may produce exit=0 on Windows.
                            # Inject an actual abnormal exit through our held process handle;
                            # the parent still closes the owned Job/tree in its finally block.
                            assert child.poll() is None
                            child.kill()
                        return child

                    setattr(postprocess, "launch_owned", launch)
                    if timer:
                        timer.start()
                    try:
                        postprocess.convert_raw(
                            "\n".join(f"synthetic {i}" for i in range(5000)),
                            fault_target,
                            cancel=fault_cancel,
                        )
                    except (ValueError, TimeoutError, Cancelled) as error:
                        expected = {
                            "parser": "PARSER_TIMEOUT",
                            "writer": "WRITER_TIMEOUT",
                            "crash": "RESOURCE_FAILURE",
                            "cancel": "CANCELLED",
                        }[fault]
                        assert expected in str(error), (fault, type(error).__name__)
                    else:
                        raise AssertionError("Frozen fault was not observed: " + fault)
                    finally:
                        if timer:
                            timer.join()
                        postprocess.PARSER_SECONDS, postprocess.WRITER_SECONDS = (
                            old_parser,
                            old_writer,
                        )
                        setattr(postprocess, "launch_owned", old_launch)
                    assert not fault_target.exists()
                    postprocess.convert_raw("next after fault", fault_target)
                    book = load_workbook(fault_target)
                    assert book["Распознанный текст"]["A1"].value == "next after fault"
                    book.close()
                    checks["frozen_worker_" + fault + "_then_next"] = "PASS"
            finally:
                postprocess.PARSER_SECONDS, postprocess.WRITER_SECONDS = old_parser, old_writer
                setattr(postprocess, "launch_owned", old_launch)
        finally:
            runtime.ROOT_OVERRIDE = old_root
            setattr(postprocess, "worker_argv", old_argv)
        log = (moved / "localocr.log").read_text("utf-8")
        assert "synthetic-evaluation-only" not in log
        checks["empty_authorization_absent_and_token_redacted"] = "PASS"
        assert not list(work.rglob(".localocr-pending-*"))
        assert not any(cwd.iterdir())
        assert all(
            not (work / key / "LocalOCR").exists()
            for key in ("APPDATA", "LOCALAPPDATA", "XDG_CONFIG_HOME")
        )
        assert not (moved / "LocalOCR-CLI/config.json").exists()
        assert not (moved / "_internal/config.json").exists()
        checks["portable_root_scoped_write_audit"] = "PASS"
    return {
        "status": "PASS",
        "platform": sys.platform,
        "version": version,
        "commit": commit,
        "checks": checks,
        "real_ocr": 0,
        "manual_gui": "PENDING",
        "server": "PENDING",
        "embedded": "NOT_BUNDLED",
        "production_release": "NO_GO",
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--root", required=True, type=Path)
    p.add_argument("--version", required=True)
    p.add_argument("--commit", required=True)
    p.add_argument("--output", required=True, type=Path)
    p.add_argument("--zip", type=Path)
    a = p.parse_args()
    started = time.time()
    if a.zip:
        archive = inspect_zip(a.zip)
        with tempfile.TemporaryDirectory(prefix="LocalOCR ZIP Русский ") as extracted:
            with zipfile.ZipFile(a.zip) as z:
                z.extractall(extracted)
                if sys.platform != "win32":
                    for executable in (
                        "LocalOCR/LocalOCR.exe",
                        "LocalOCR/LocalOCR-CLI/LocalOCR-CLI.exe",
                    ):
                        (Path(extracted) / executable).chmod(0o700)
            result = qualify(Path(extracted) / "LocalOCR", a.version, a.commit)
        result["archive"] = archive
    else:
        result = qualify(a.root.resolve(), a.version, a.commit)
    result["seconds"] = round(time.time() - started, 2)
    a.output.parent.mkdir(parents=True, exist_ok=True)
    a.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
