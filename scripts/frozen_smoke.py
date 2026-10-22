"""Real frozen CLI -> local HTTP mock -> XLSX roundtrip (no external network)."""

import argparse
import json
import os
import subprocess
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw


def isolated_environment():
    env = dict(os.environ)
    for name in ("PYTHONHOME", "PYTHONPATH", "VIRTUAL_ENV"):
        env.pop(name, None)
    env["PATH"] = (
        str(Path(env.get("SystemRoot", "C:/Windows")) / "System32")
        if os.name == "nt"
        else "/usr/bin:/bin"
    )
    return env


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--exe", required=True, type=Path)
    parser.add_argument("--version", required=True)
    parser.add_argument("--commit")
    args = parser.parse_args()
    exe = args.exe.resolve()
    version = subprocess.check_output(
        [str(exe), "--version"], text=True, encoding="utf-8", env=isolated_environment()
    ).strip()
    if version != args.version:
        raise RuntimeError(f"Frozen version mismatch: {version}")

    if args.commit:
        commit = subprocess.check_output(
            [str(exe), "--build-commit"], text=True, encoding="utf-8", env=isolated_environment()
        ).strip()
        if commit != args.commit:
            raise RuntimeError(f"Frozen commit mismatch: {commit}")

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, format: str, *args):
            pass

        def do_GET(self):
            self.respond({"models": [{"name": "glm-ocr:latest"}]})

        def do_POST(self):
            payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            if payload["stream"] or not payload["images"]:
                raise RuntimeError("Incorrect payload")
            self.respond(
                {
                    "response": "<table><tr><th>Код</th><th>Значение</th></tr><tr><td>0012</td><td>=1+1</td></tr></table>",
                    "done": True,
                }
            )

        def respond(self, data):
            self.send_response(200)
            self.end_headers()
            self.wfile.write(json.dumps(data).encode())

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with tempfile.TemporaryDirectory(prefix="LocalOCR smoke ") as name:
            folder = Path(name)
            image = folder / "Таблица.png"
            sample = Image.new("RGB", (200, 100), "white")
            ImageDraw.Draw(sample).text((10, 10), "table rows", fill="black")
            sample.save(image)
            command = [
                str(exe),
                "convert",
                str(image),
                "--ollama-url",
                f"http://127.0.0.1:{server.server_port}",
            ]
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=60,
                env=isolated_environment(),
            )
            # Unruled but visibly structured input must honestly be REVIEW, exit 1.
            if result.returncode != 1:
                raise RuntimeError(f"Frozen conversion failed: {result.stdout} {result.stderr}")
            output = folder / "results" / "Таблица.xlsx"
            if not output.is_file():
                raise RuntimeError(
                    f"Frozen conversion did not publish its REVIEW workbook: {result.stdout} {result.stderr}"
                )
            book = load_workbook(output)
            assert book.worksheets[0]["A2"].value == "0012"
            assert book.worksheets[0]["B2"].data_type == "s"
            assert "Проверка" in book.sheetnames
            book.close()
            print("Frozen CLI, version, Unicode paths and XLSX: PASS")
    finally:
        server.shutdown()
        server.server_close()


if __name__ == "__main__":
    main()
